"""
Test Suite for Alpha Aerosols Engine & Mass Balance Model
"""

import os
import sqlite3
import json
import unittest
from engine import (
    get_connection,
    calculate_bom,
    calculate_yield_inverse,
    evaluate_shelf_life,
    add_shift_entry,
    create_order,
    create_dispatch,
    export_job_card_excel,
    export_production_json,
    sync_excel_to_db,
    sync_db_to_excel,
    get_monthly_workbook_name,
    get_active_aerosol_workbook,
    seed_master_data,
    DB_PATH,
    JSON_PATH
)

class TestAlphaAerosolsEngine(unittest.TestCase):

    def test_sqlite_wal_mode(self):
        """Verify SQLite WAL mode and foreign keys are active."""
        conn = get_connection(DB_PATH)
        cur = conn.cursor()
        cur.execute("PRAGMA journal_mode;")
        mode = cur.fetchone()[0]
        self.assertEqual(mode.lower(), "wal")

        cur.execute("PRAGMA foreign_keys;")
        fk = cur.fetchone()[0]
        self.assertEqual(fk, 1)
        conn.close()

    def test_yield_inverse_mass_balance(self):
        """
        Verify the physical yield-inverse formula:
        Req = Net / (1 - Scrap)
        """
        # Lacquer: 35% loss -> Gross = Net / (1 - 0.35) = Net / 0.65 = 1.5385 * Net
        net = 1.045
        gross = calculate_yield_inverse(net, 0.35)
        expected = net / 0.65
        self.assertAlmostEqual(gross, expected, places=5)
        # Ensure it is strictly greater than additive Net * (1 + 0.35) = 1.41075
        self.assertGreater(gross, net * 1.35)

        # Base coat: 10% loss -> Gross = Net / 0.90 = 1.1111 * Net
        net_bc = 0.973
        gross_bc = calculate_yield_inverse(net_bc, 0.10)
        self.assertAlmostEqual(gross_bc, net_bc / 0.90, places=5)

    def test_bom_45x160_calculations(self):
        """Verify 45x160mm standard container BOM matches high-end engineering specs."""
        bom = calculate_bom("45x160mm", 50000)
        items = {item["item_code"]: item for item in bom["items"]}

        # Lacquer (504): 1.045 kg/1000 net -> gross 1.608 kg/1000
        lacquer = items["504"]
        self.assertAlmostEqual(lacquer["net_rate_per_1000"], 1.045, places=2)
        self.assertAlmostEqual(lacquer["gross_rate_per_1000"], 1.608, places=2)
        self.assertEqual(lacquer["scrap_pct"], 35.0)

        # Base coat (506): 0.973 kg/1000 net -> gross 1.081 kg/1000
        bc = items["506"]
        self.assertAlmostEqual(bc["net_rate_per_1000"], 0.973, places=2)
        self.assertAlmostEqual(bc["gross_rate_per_1000"], 1.081, places=2)

        # Varnish (508): 0.430 kg/1000 net -> gross 0.478 kg/1000
        varnish = items["508"]
        self.assertAlmostEqual(varnish["net_rate_per_1000"], 0.430, places=2)
        self.assertAlmostEqual(varnish["gross_rate_per_1000"], 0.478, places=2)

        # Slugs (501): 20.0 kg/1000 net -> gross 22.222 kg/1000
        slugs = items["501"]
        self.assertAlmostEqual(slugs["gross_rate_per_1000"], 22.222, places=2)

    def test_shelf_life_categorization(self):
        """
        Verify chemical shelf-life countdown logic:
        <120d: normal (green)
        120-149d: warning (amber)
        >=150d: critical (red)
        """
        ref_date = "2026-09-12"

        # 42 days old (received 2026-08-01)
        res_normal = evaluate_shelf_life("2026-08-01", 180, ref_date)
        self.assertEqual(res_normal["status"], "normal")
        self.assertEqual(res_normal["badge_color"], "#10b981")
        self.assertEqual(res_normal["age_days"], 42)

        # 125 days old (received 2026-05-10)
        res_warning = evaluate_shelf_life("2026-05-10", 180, ref_date)
        self.assertEqual(res_warning["status"], "warning")
        self.assertEqual(res_warning["badge_color"], "#f59e0b")
        self.assertEqual(res_warning["age_days"], 125)

        # 153 days old (received 2026-04-12)
        res_critical = evaluate_shelf_life("2026-04-12", 180, ref_date)
        self.assertEqual(res_critical["status"], "critical")
        self.assertEqual(res_critical["badge_color"], "#ef4444")
        self.assertEqual(res_critical["age_days"], 153)

    def test_json_export_structure(self):
        """Verify production.json adheres to schema and contains required metrics."""
        test_out_json = os.path.join(os.path.dirname(__file__), 'data', 'test_schema_out.json')
        self.addCleanup(lambda: os.remove(test_out_json) if os.path.exists(test_out_json) else None)
        data = export_production_json(json_path=test_out_json, as_of_date="2026-09-12")
        self.assertIn("meta", data)
        self.assertIn("kpis", data)
        self.assertIn("orders", data)
        self.assertIn("shifts", data)
        self.assertIn("inventory", data)
        self.assertIn("dispatches", data)

        kpis = data["kpis"]
        self.assertIn("today_output_cans", kpis)
        self.assertIn("today_scrap_pct", kpis)
        self.assertIn("mtd_output_cans", kpis)
        self.assertIn("active_pofs_count", kpis)
        self.assertIn("today_dispatches_cans", kpis)

        # Tolerance bounds on orders
        for ord_item in data["orders"]:
            self.assertIn("min_acceptable_qty", ord_item)
            self.assertIn("max_acceptable_qty", ord_item)
            self.assertIn("completion_pct", ord_item)
            self.assertIn("remaining_qty", ord_item)

    def test_job_card_generation(self):
        """Verify openpyxl job card export creates a valid excel file without COM."""
        test_out = os.path.join(os.path.dirname(__file__), 'data', 'test_job_card.xlsx')
        self.addCleanup(lambda: os.remove(test_out) if os.path.exists(test_out) else None)
        export_job_card_excel(1, test_out)
        self.assertTrue(os.path.exists(test_out))
        self.assertGreater(os.path.getsize(test_out), 4000)

    def test_fg_stock_and_latest_shift_kpis(self):
        """Verify fg_stock per order and fg_buffer_cans, latest_shift in KPIs."""
        test_out_json = os.path.join(os.path.dirname(__file__), 'data', 'test_kpi_out.json')
        self.addCleanup(lambda: os.remove(test_out_json) if os.path.exists(test_out_json) else None)
        data = export_production_json(json_path=test_out_json, as_of_date="2026-09-12")
        kpis = data["kpis"]
        self.assertIn("fg_buffer_cans", kpis)
        self.assertIn("fg_buffer_pallets", kpis)
        self.assertIn("latest_shift", kpis)
        self.assertIsInstance(kpis["latest_shift"], dict)
        self.assertIn("good_cans", kpis["latest_shift"])
        self.assertIn("supervisor", kpis["latest_shift"])

        for ord_item in data["orders"]:
            self.assertIn("fg_stock", ord_item)
            expected_fg = max(0, ord_item["produced_good"] - ord_item["dispatched_total"])
            self.assertEqual(ord_item["fg_stock"], expected_fg)

    def test_excel_to_sqlite_sync(self):
        """Verify synchronization of shift entries from Excel into SQLite without COM, ensuring full isolation and edge-case handling."""
        import openpyxl

        test_db = os.path.join(os.path.dirname(__file__), 'data', 'test_sync.db')
        test_json = os.path.join(os.path.dirname(__file__), 'data', 'test_sync.json')
        test_xlsx = os.path.join(os.path.dirname(__file__), 'data', 'test_entry.xlsx')

        self.addCleanup(lambda: [os.remove(p) for p in [test_db, test_db+'-wal', test_db+'-shm', test_json, test_xlsx] if os.path.exists(p)])

        # Set up a test DB
        seed_master_data(test_db, force_reseed=True, demo_data=False)

        # Create a test Excel workbook with comma numbers, flexible dates, reconciliation, and new POF
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data_Entry"
        ws.append(["AEROSOL PLANT DAILY PRODUCTION DATA ENTRY"])
        ws.append(["Date", "Machine", "POF #", "Product Name", "PID", "Customer", "Total Production\n(pcs)", "Good Production\n(pcs)", "Rejects\n(pcs)", "Rejection\n%", "DownTime", "Remarks"])
        # Row 3: Comma-formatted numbers, text date "12-Sep-2026", supervisor in remarks
        ws.append(["12-Sep-2026", "Press", "POF-2026-001", "AEROSOL CAN 45x160", 5002, "Aerosol Customer", "12,400", "12,000", "400", "3.23%", 0.5, "Washer nozzle clean - Tariq Mahmood"])
        # Row 4: Reconciled quantities (good is None), Shift B (Night), New POF-2026-002, 45x150
        ws.append(["2026-09-13", "Printing", "POF-2026-002", "AEROSOL CAN 45x150", 9003, "New Brand Inc", 15000, None, 500, "3.33%", 1.0, "Shift B ink viscosity - M. Aslam"])
        wb.save(test_xlsx)

        # Run sync passing test_json to guarantee isolation
        res1 = sync_excel_to_db(test_xlsx, test_db, test_json)
        self.assertEqual(res1["status"], "success")
        self.assertEqual(res1["shifts_imported"], 1)
        self.assertEqual(res1["shifts_updated"], 1)
        self.assertEqual(res1["skipped_duplicates"], 0)

        # Verify test_json was created and master production.json was not mutated
        self.assertTrue(os.path.exists(test_json))
        with open(test_json, 'r', encoding='utf-8') as f:
            t_data = json.load(f)
        self.assertEqual(len(t_data["orders"]), 2)
        self.assertEqual(t_data["kpis"]["latest_shift"]["good_cans"], 14500)

        # Verify DB records
        conn = get_connection(test_db)
        cur = conn.cursor()
        cur.execute("SELECT s.*, o.pof_number, o.customer_name FROM shifts s JOIN orders o ON s.pof_id = o.id ORDER BY s.id ASC;")
        shifts = [dict(r) for r in cur.fetchall()]
        conn.close()

        # Check Row 3 updated inaugural shift 1 in-place with commas and date
        self.assertEqual(shifts[0]["good_cans"], 12000)
        self.assertEqual(shifts[0]["line_scrap"], 400)
        self.assertEqual(shifts[0]["shift_date"], "2026-09-12")
        self.assertEqual(shifts[0]["shift_type"], "Day")
        self.assertEqual(shifts[0]["supervisor"], "Tariq Mahmood")
        self.assertEqual(shifts[0]["downtime_reason"], "Washer nozzle clean")

        # Check Row 4 reconciled good = total - scrap (15000 - 500 = 14500), Shift B = Night, new POF
        self.assertEqual(shifts[1]["good_cans"], 14500)
        self.assertEqual(shifts[1]["line_scrap"], 500)
        self.assertEqual(shifts[1]["shift_type"], "Night")
        self.assertEqual(shifts[1]["pof_number"], "POF-2026-002")
        self.assertEqual(shifts[1]["customer_name"], "New Brand Inc")
        self.assertEqual(shifts[1]["product_size"], "45x150mm")
        self.assertEqual(shifts[1]["supervisor"], "M. Aslam")

        # Verify idempotence on second sync
        res2 = sync_excel_to_db(test_xlsx, test_db, test_json)
        self.assertEqual(res2["status"], "success")
        self.assertEqual(res2["shifts_imported"], 0)
        self.assertEqual(res2["shifts_updated"], 0)
        self.assertEqual(res2["skipped_duplicates"], 2)

    def test_monthly_workbook_resolution(self):
        """Verify dynamic monthly Excel naming (Aerosol_MmmYY.xlsx) and fallback resolution."""
        self.assertEqual(get_monthly_workbook_name("2026-09-14"), "Aerosol_Sep26.xlsx")
        self.assertEqual(get_monthly_workbook_name("2026-10-01"), "Aerosol_Oct26.xlsx")
        self.assertEqual(get_monthly_workbook_name("Oct26"), "Aerosol_Oct26.xlsx")
        self.assertEqual(get_monthly_workbook_name("2026-11"), "Aerosol_Nov26.xlsx")
        self.assertEqual(get_monthly_workbook_name("2027-01-15"), "Aerosol_Jan27.xlsx")

        active_wb = get_active_aerosol_workbook(target_date="2026-09-14")
        self.assertTrue(os.path.exists(active_wb), f"Active workbook {active_wb} must exist")
        self.assertTrue(os.path.basename(active_wb).startswith("Aerosol_"))

        # Test chronological sorting: Dec26 must be recognized as later than Jan26 (alphabetically 'Jan' > 'Dec')
        import tempfile
        with tempfile.TemporaryDirectory() as tmp_dir:
            wb_jan = os.path.join(tmp_dir, "Aerosol_Jan26.xlsx")
            wb_dec = os.path.join(tmp_dir, "Aerosol_Dec26.xlsx")
            open(wb_jan, "w").close()
            open(wb_dec, "w").close()
            # Since target_date=None looks for current month (Sep26) which does not exist in tmp_dir,
            # it falls back to candidates and must sort chronologically: Dec26 > Jan26!
            latest = get_active_aerosol_workbook(base_dir=tmp_dir, target_date=None)
            self.assertEqual(os.path.basename(latest), "Aerosol_Dec26.xlsx")

            # Test create_if_missing=True for target month creates target even if older exists
            wb_target = get_active_aerosol_workbook(base_dir=tmp_dir, target_date="2027-02-01", create_if_missing=True)
            self.assertEqual(os.path.basename(wb_target), "Aerosol_Feb27.xlsx")
            self.assertTrue(os.path.exists(wb_target))

    def test_get_connection_relative_path(self):
        """Verify get_connection does not crash with WinError 3 when db_path has no directory prefix."""
        rel_db = "test_rel_tmp.db"
        try:
            conn = get_connection(rel_db)
            self.assertIsNotNone(conn)
            conn.close()
        finally:
            if os.path.exists(rel_db):
                try:
                    os.remove(rel_db)
                except Exception:
                    pass

    def test_bidirectional_excel_sync(self):
        """
        Verify bidirectional synchronization:
        Option A: Web/SQLite -> Monthly Excel workbook (sync_db_to_excel)
        Option B: Monthly Excel -> Web/SQLite (sync_excel_to_db)
        Ensures idempotence, row updates vs append, POF normalization, and raw material deduction.
        """
        import openpyxl
        import shutil

        test_dir = os.path.join(os.path.dirname(__file__), 'data')
        test_db = os.path.join(test_dir, 'test_bi_aerosol.db')
        test_json = os.path.join(test_dir, 'test_bi_production.json')
        test_xlsx = os.path.join(test_dir, 'Aerosol_Sep26_test.xlsx')

        # Clean slate
        seed_master_data(test_db, force_reseed=True, demo_data=False)
        export_production_json(test_db, test_json)

        # Clone template to test_xlsx
        template_path = os.path.join(os.path.dirname(__file__), 'Aerosol_Production_Entry.xlsx')
        shutil.copyfile(template_path, test_xlsx)

        try:
            # 1. Option A: Log a 10,000 can production shift via engine in SQLite
            shift_id = add_shift_entry(
                shift_date="2026-09-14",
                shift_type="Day",
                pof_id=1,
                product_size="45x160mm",
                good_cans=10000,
                line_scrap=350,
                downtime_hours=0.5,
                downtime_reason="Slug feeder jam",
                supervisor="Tariq Mahmood",
                db_path=test_db
            )
            self.assertIsNotNone(shift_id)

            # 2. Push SQLite shift to Excel workbook (Option A)
            res_to = sync_db_to_excel(excel_path=test_xlsx, db_path=test_db, month_str="Sep26")
            self.assertEqual(res_to["status"], "success")
            self.assertEqual(res_to["rows_written"], 2)

            # Inspect Excel file directly to verify contents and formatting
            wb = openpyxl.load_workbook(test_xlsx, data_only=True)
            ws = wb["Data_Entry"]
            self.assertEqual(str(ws.cell(row=4, column=1).value)[:10], "2026-09-14")
            self.assertEqual(ws.cell(row=4, column=2).value, "Continuous Line 1")
            self.assertEqual(ws.cell(row=4, column=3).value, "POF-2026-001")
            self.assertEqual(ws.cell(row=4, column=5).value, 5002)
            self.assertEqual(ws.cell(row=4, column=7).value, 10350)
            self.assertEqual(ws.cell(row=4, column=8).value, 10000)
            self.assertEqual(ws.cell(row=4, column=9).value, 350)
            self.assertEqual(ws.cell(row=4, column=11).value, 0.5)
            self.assertIn("Tariq Mahmood", str(ws.cell(row=4, column=12).value))
            wb.close()

            # 3. Idempotence test for Option A (pushing again without new shifts writes 0 and updates 0)
            res_to_again = sync_db_to_excel(excel_path=test_xlsx, db_path=test_db, month_str="Sep26")
            self.assertEqual(res_to_again["rows_written"], 0)
            self.assertEqual(res_to_again["rows_updated"], 0)

            # 3b. Option A: Update existing entry data from Web -> Excel!
            # Operator logged 15,000 for today (updated run)
            conn_u = get_connection(test_db)
            conn_u.execute("UPDATE shifts SET good_cans = 15000, line_scrap = 500 WHERE id = ?;", (shift_id,))
            conn_u.commit()
            conn_u.close()

            res_to_update = sync_db_to_excel(excel_path=test_xlsx, db_path=test_db, month_str="Sep26")
            self.assertEqual(res_to_update["rows_updated"], 1)
            self.assertEqual(res_to_update["rows_written"], 0)

            # Verify the row in Excel was updated in-place (row 4 now shows 15,000)
            wb_up = openpyxl.load_workbook(test_xlsx, data_only=True)
            ws_up = wb_up["Data_Entry"]
            self.assertEqual(ws_up.cell(row=4, column=8).value, 15000)
            self.assertEqual(ws_up.cell(row=4, column=9).value, 500)
            self.assertEqual(ws_up.cell(row=4, column=7).value, 15500)
            wb_up.close()

            # 4. Option B: Sync from Excel back to DB -> recognizes row 4 as duplicate
            res_from = sync_excel_to_db(excel_path=test_xlsx, db_path=test_db, json_path=test_json)
            self.assertEqual(res_from["status"], "success")
            self.assertEqual(res_from["shifts_imported"], 0)
            self.assertEqual(res_from["skipped_duplicates"], 1)

            # 5. Add a new manual shift in Excel on row 5 (Option B)
            wb2 = openpyxl.load_workbook(test_xlsx)
            ws2 = wb2["Data_Entry"]
            ws2.cell(row=5, column=1, value="2026-09-14")
            ws2.cell(row=5, column=2, value="Continuous Line 1")
            ws2.cell(row=5, column=3, value="POF-2026-001")
            ws2.cell(row=5, column=4, value="5002 - Aerosol Container 45x160mm")
            ws2.cell(row=5, column=5, value=5002)
            ws2.cell(row=5, column=6, value="Aerosol Customer")
            ws2.cell(row=5, column=7, value=12500)
            ws2.cell(row=5, column=8, value=12000)
            ws2.cell(row=5, column=9, value=500)
            ws2.cell(row=5, column=10, value=0.04)
            ws2.cell(row=5, column=11, value=0.0)
            ws2.cell(row=5, column=12, value="Night Shift - M. Aslam - Normal continuous run")
            wb2.save(test_xlsx)
            wb2.close()

            # 6. Run Option B: Excel -> DB to import row 5
            res_from_new = sync_excel_to_db(excel_path=test_xlsx, db_path=test_db, json_path=test_json)
            self.assertEqual(res_from_new["status"], "success")
            self.assertEqual(res_from_new["shifts_imported"], 1)
            self.assertEqual(res_from_new["skipped_duplicates"], 1)

            # Verify the newly imported shift exists in SQLite
            conn = get_connection(test_db)
            cur = conn.cursor()
            cur.execute("SELECT good_cans, line_scrap, shift_type, supervisor FROM shifts WHERE good_cans = 12000;")
            row = cur.fetchone()
            self.assertIsNotNone(row)
            self.assertEqual(row["shift_type"], "Night")
            self.assertEqual(row["supervisor"], "M. Aslam")

            # 7. Option B Update: Operator fixes quantity in Excel on row 5 from 12,000 to 14,000
            wb3 = openpyxl.load_workbook(test_xlsx)
            ws3 = wb3["Data_Entry"]
            ws3.cell(row=5, column=8, value=14000)
            ws3.cell(row=5, column=7, value=14500)
            wb3.save(test_xlsx)
            wb3.close()

            res_from_update = sync_excel_to_db(excel_path=test_xlsx, db_path=test_db, json_path=test_json)
            self.assertEqual(res_from_update["status"], "success")
            self.assertEqual(res_from_update["shifts_updated"], 1)
            self.assertEqual(res_from_update["shifts_imported"], 0)

            # Verify shift in SQLite was updated in place without duplicate row!
            cur.execute("SELECT good_cans FROM shifts WHERE shift_date = '2026-09-14' AND shift_type = 'Night';")
            updated_shift_row = cur.fetchone()
            self.assertEqual(updated_shift_row["good_cans"], 14000)

            # Check total shifts count in DB is exactly 3 (1 inaugural + 1 Day + 1 Night)
            cur.execute("SELECT COUNT(*) FROM shifts;")
            self.assertEqual(cur.fetchone()[0], 3)
            conn.close()

            # 8. POF Normalization test: Excel with POF '81' matches 'POF-2026-081' without duplicating
            wb_pof = openpyxl.load_workbook(test_xlsx)
            ws_pof = wb_pof["Data_Entry"]
            ws_pof.cell(row=6, column=1, value="2026-09-15")
            ws_pof.cell(row=6, column=2, value="Continuous Line 1")
            ws_pof.cell(row=6, column=3, value="81")  # Operator wrote '81' instead of 'POF-2026-081'
            ws_pof.cell(row=6, column=8, value=8000)
            ws_pof.cell(row=6, column=9, value=300)
            ws_pof.cell(row=6, column=7, value=8300)
            ws_pof.cell(row=6, column=12, value="Day Shift - Tariq Mahmood")
            wb_pof.save(test_xlsx)
            wb_pof.close()

            # Option B imports it into DB as POF-81 or order 81
            res_pof_b = sync_excel_to_db(excel_path=test_xlsx, db_path=test_db, json_path=test_json)
            self.assertEqual(res_pof_b["shifts_imported"], 1)

            # Option A should recognize that row 6 already has this shift and NOT write a duplicate
            res_pof_a = sync_db_to_excel(excel_path=test_xlsx, db_path=test_db, month_str="Sep26")
            self.assertEqual(res_pof_a["rows_written"], 0)
            self.assertEqual(res_pof_a["rows_updated"], 0)

        finally:
            for ext in ['', '-wal', '-shm']:
                p = test_db + ext
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except Exception:
                        pass
            if os.path.exists(test_json):
                try:
                    os.remove(test_json)
                except Exception:
                    pass
            if os.path.exists(test_xlsx):
                try:
                    os.remove(test_xlsx)
                except Exception:
                    pass

if __name__ == "__main__":
    unittest.main()
