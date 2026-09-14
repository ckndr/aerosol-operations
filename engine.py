"""
Alpha Aerosols (Kot Abdul Malik) — Production & Operations Database Engine
Single Source of Truth: aerosol.db (SQLite with Write-Ahead Logging WAL mode)
Enforces Physical Yield-Inverse Mass Balance: Req = Net / (1 - Scrap)
Zero Windows COM Automation: Pure Python + SQLite + openpyxl
"""

import os
import sqlite3
import math
import json
import glob
import re
import shutil
from datetime import datetime, date, timedelta
from typing import Dict, List, Any, Optional, Union

DB_PATH = os.path.join(os.path.dirname(__file__), 'data', 'aerosol.db')
JSON_PATH = os.path.join(os.path.dirname(__file__), 'data', 'production.json')

# Production line capacity constants
CANS_PER_SHIFT_NOMINAL = 30000  # Nominal target output per 12h shift
CAN_SIZES = {
    "45x160mm": {
        "diameter_mm": 45.0,
        "height_mm": 160.0,
        "wall_thickness_mm": 0.40,
        "slug_weight_g": 20.0,
        "outer_surface_m2": math.pi * 0.045 * 0.160, # 0.022619 m²
        # Inner: dia = 45 - 2*0.40 = 44.20 mm
        "inner_surface_m2": (math.pi * 0.0442 * 0.160) + (math.pi * (0.0221 ** 2)), # 0.023751 m²
    },
    "45x150mm": {
        "diameter_mm": 45.0,
        "height_mm": 150.0,
        "wall_thickness_mm": 0.35,
        "slug_weight_g": 19.0,
        "outer_surface_m2": math.pi * 0.045 * 0.150, # 0.021206 m²
        # Inner: dia = 45 - 2*0.35 = 44.30 mm
        "inner_surface_m2": (math.pi * 0.0443 * 0.150) + (math.pi * (0.02215 ** 2)), # 0.022417 m²
    }
}

# Locked Scrap Factors (Transfer Losses)
SCRAP_FACTORS = {
    "lacquer": 0.35,     # 35% airless spray transfer loss
    "base_coat": 0.10,   # 10% roller coater loss
    "varnish": 0.10,     # 10% roller coater loss
    "ink": 0.10,         # 10% dry offset ink loss
    "slugs": 0.10,       # 10% overall mechanical line scrap
    "lubricant": 0.10,   # 10% process loss
    "cleaner": 0.10,     # 10% wash bath loss
}

def get_connection(db_path: Optional[str] = None) -> sqlite3.Connection:
    """Returns a SQLite connection configured with WAL mode and foreign keys enabled."""
    if db_path is None:
        db_path = DB_PATH
    db_dir = os.path.dirname(db_path)
    if db_dir:
        os.makedirs(db_dir, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = WAL;")
    conn.execute("PRAGMA foreign_keys = ON;")
    conn.execute("PRAGMA synchronous = NORMAL;")
    return conn

def init_db(db_path: Optional[str] = None):
    """Initializes the database schema."""
    if db_path is None:
        db_path = DB_PATH
    conn = get_connection(db_path)
    with conn:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pof_number TEXT UNIQUE NOT NULL,
            customer_name TEXT NOT NULL,
            product_name TEXT NOT NULL,
            product_size TEXT NOT NULL,
            order_qty INTEGER NOT NULL,
            tolerance_pct REAL NOT NULL DEFAULT 0.05,
            order_date TEXT NOT NULL,
            due_date TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'In Production',
            artwork_ref TEXT,
            notes TEXT,
            created_at TEXT NOT NULL
        );
        """)

        conn.execute("""
        CREATE TABLE IF NOT EXISTS shifts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shift_date TEXT NOT NULL,
            shift_type TEXT NOT NULL,
            pof_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE RESTRICT,
            product_size TEXT NOT NULL,
            good_cans INTEGER NOT NULL,
            line_scrap INTEGER NOT NULL,
            downtime_hours REAL NOT NULL DEFAULT 0.0,
            downtime_reason TEXT,
            supervisor TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        """)

        conn.execute("""
        CREATE TABLE IF NOT EXISTS inventory (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_code TEXT UNIQUE NOT NULL,
            category TEXT NOT NULL,
            item_name TEXT NOT NULL,
            uom TEXT NOT NULL,
            balance_qty REAL NOT NULL,
            min_stock_level REAL NOT NULL DEFAULT 0,
            unit_cost_pkr REAL DEFAULT 0,
            batch_number TEXT,
            received_date TEXT NOT NULL,
            shelf_life_days INTEGER NOT NULL,
            storage_condition TEXT NOT NULL,
            location TEXT DEFAULT 'Kot Abdul Malik Central Stores',
            updated_at TEXT NOT NULL
        );
        """)

        conn.execute("""
        CREATE TABLE IF NOT EXISTS dispatches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            challan_number TEXT UNIQUE NOT NULL,
            dispatch_date TEXT NOT NULL,
            pof_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE RESTRICT,
            dispatched_cans INTEGER NOT NULL,
            carton_count INTEGER NOT NULL,
            pallet_count INTEGER NOT NULL,
            vehicle_number TEXT NOT NULL,
            driver_name TEXT,
            receiver_party TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Dispatched',
            created_at TEXT NOT NULL
        );
        """)

        conn.execute("""
        CREATE TABLE IF NOT EXISTS bom_standards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_size TEXT NOT NULL,
            material_group TEXT NOT NULL,
            item_code TEXT NOT NULL,
            item_name TEXT NOT NULL,
            uom TEXT NOT NULL,
            net_rate_per_1000 REAL NOT NULL,
            scrap_rate REAL NOT NULL,
            gross_rate_per_1000 REAL NOT NULL,
            calculation_rule TEXT NOT NULL
        );
        """)
    conn.close()

def calculate_yield_inverse(net_rate: float, scrap_factor: float) -> float:
    """
    Physical Yield-Inverse Mass Balance:
    Gross = Net / (1 - Scrap)
    Avoids Tubex additive error of Net * (1 + Scrap).
    """
    if scrap_factor >= 1.0:
        raise ValueError("Scrap factor must be less than 1.0")
    return net_rate / (1.0 - scrap_factor)

def calculate_bom(product_size: str, can_qty: int, options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """
    Computes rigorous mass-balance requirements for any order quantity.
    Supports 45x160mm (standard) and 45x150mm.
    Enforces yield-inverse physics: Req = Net / (1 - Scrap).
    """
    size_spec = CAN_SIZES.get(product_size, CAN_SIZES["45x160mm"])
    thousand_units = can_qty / 1000.0

    # Net Specs based on geometry & High-End TDS values
    # Internal Lacquer (SCHEKOSOL INT PROT Gold - 44 g/m² vs Beige - 48 g/m²)
    lacquer_type = (options and options.get("lacquer_type")) or "gold"
    if str(lacquer_type).lower() == "beige":
        lacquer_item_code = "505"
        lacquer_item_name = "SCHEKOSOL INT BEIGE (400 4 902)"
        net_lacquer_per_1000 = 1.140 # kg/1000 cans
        lacquer_formula = "Inner Area (m²) x 48g/m² / (1 - 0.35) [Transfer efficiency 65%]"
    else:
        lacquer_item_code = "504"
        lacquer_item_name = "SCHEKOSOL INT PROT Gold (400 9 901)"
        net_lacquer_per_1000 = size_spec["inner_surface_m2"] * 44.0 # kg/1000 cans
        lacquer_formula = "Inner Area (m²) x 44g/m² / (1 - 0.35) [Transfer efficiency 65%]"
    gross_lacquer_per_1000 = calculate_yield_inverse(net_lacquer_per_1000, SCRAP_FACTORS["lacquer"])

    # Base Coat (SCHEKOSOL WH BC - 43 g/m² vs CLEAR BC - 31 g/m²)
    base_coat_type = (options and options.get("base_coat_type")) or "white"
    if str(base_coat_type).lower() == "clear":
        base_item_code = "507"
        base_item_name = "SCHEKOSOL CLEAR BC (422 9 918)"
        net_base_per_1000 = 0.701 # kg/1000 cans
        base_formula = "Outer Area (m²) x 31g/m² / (1 - 0.10) [Transfer efficiency 90%]"
    else:
        base_item_code = "506"
        base_item_name = "SCHEKOSOL WH BC (422 0 903)"
        net_base_per_1000 = size_spec["outer_surface_m2"] * 43.0 # kg/1000 cans
        base_formula = "Outer Area (m²) x 43g/m² / (1 - 0.10) [Transfer efficiency 90%]"
    gross_base_per_1000 = calculate_yield_inverse(net_base_per_1000, SCRAP_FACTORS["base_coat"])

    # Varnish (SCHEKOSOL OPV GLOSSY - 19 g/m² vs SILKMATT - 17 g/m²)
    varnish_type = (options and options.get("varnish_type")) or "glossy"
    if str(varnish_type).lower() == "silkmatt":
        varnish_item_code = "509"
        varnish_item_name = "SCHEKOSOL OPV SILKMATT (422 9 903)"
        net_varnish_per_1000 = 0.385 # kg/1000 cans
        varnish_formula = "Outer Area (m²) x 17g/m² / (1 - 0.10) [Transfer efficiency 90%]"
    else:
        varnish_item_code = "508"
        varnish_item_name = "SCHEKOSOL OPV GLOSSY (422 9 900)"
        net_varnish_per_1000 = size_spec["outer_surface_m2"] * 19.0 # kg/1000 cans
        varnish_formula = "Outer Area (m²) x 19g/m² / (1 - 0.10) [Transfer efficiency 90%]"
    gross_varnish_per_1000 = calculate_yield_inverse(net_varnish_per_1000, SCRAP_FACTORS["varnish"])

    # Slugs
    net_slug_per_1000 = size_spec["slug_weight_g"] # 20.0 kg/1000 for 160mm, 19.0 for 150mm
    gross_slug_per_1000 = calculate_yield_inverse(net_slug_per_1000, SCRAP_FACTORS["slugs"])

    # Extrusion Lubricant (105g per 100kg slugs)
    net_lubricant_per_1000 = (net_slug_per_1000 * 0.105) / 100.0 # 0.021 kg/1000
    gross_lubricant_per_1000 = calculate_yield_inverse(net_lubricant_per_1000, SCRAP_FACTORS["lubricant"])

    # Washer Cleaner (5.0 kg/1000 cans empirical startup buffer)
    net_cleaner_per_1000 = 5.0
    gross_cleaner_per_1000 = calculate_yield_inverse(net_cleaner_per_1000, SCRAP_FACTORS["cleaner"])

    # Printing Inks (Locked at 0.280 kg/1000 net per color station)
    num_colors = (options and options.get("num_colors")) or 4
    net_ink_per_color_per_1000 = 0.280
    gross_ink_per_color_per_1000 = calculate_yield_inverse(net_ink_per_color_per_1000, SCRAP_FACTORS["ink"])
    total_net_ink_per_1000 = net_ink_per_color_per_1000 * num_colors
    total_gross_ink_per_1000 = gross_ink_per_color_per_1000 * num_colors

    items = [
        {
            "item_code": "501",
            "material_group": "Aluminum Slugs",
            "item_name": f"{size_spec['diameter_mm']:.0f}mm Aluminum Slug (Al99.7%)",
            "uom": "kg",
            "net_rate_per_1000": round(net_slug_per_1000, 3),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(gross_slug_per_1000, 3),
            "net_total": round(net_slug_per_1000 * thousand_units, 2),
            "gross_total": round(gross_slug_per_1000 * thousand_units, 2),
            "formula": "Slug wt (g) x 1,000 / (1 - 0.10)"
        },
        {
            "item_code": "502",
            "material_group": "Extrusion Lubricant",
            "item_name": "SAPILUB LUBRIMET GR8",
            "uom": "kg",
            "net_rate_per_1000": round(net_lubricant_per_1000, 4),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(gross_lubricant_per_1000, 4),
            "net_total": round(net_lubricant_per_1000 * thousand_units, 3),
            "gross_total": round(gross_lubricant_per_1000 * thousand_units, 3),
            "formula": "105g / 100kg Slugs / (1 - 0.10)"
        },
        {
            "item_code": "503",
            "material_group": "Alkaline Cleaner",
            "item_name": "SAPILUB ALULIQUID 13",
            "uom": "kg",
            "net_rate_per_1000": round(net_cleaner_per_1000, 3),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(gross_cleaner_per_1000, 3),
            "net_total": round(net_cleaner_per_1000 * thousand_units, 2),
            "gross_total": round(gross_cleaner_per_1000 * thousand_units, 2),
            "formula": "Empirical 5.0 kg / 1,000 / (1 - 0.10)"
        },
        {
            "item_code": lacquer_item_code,
            "material_group": "Internal Lacquer",
            "item_name": lacquer_item_name,
            "uom": "kg",
            "net_rate_per_1000": round(net_lacquer_per_1000, 3),
            "scrap_pct": 35.0,
            "gross_rate_per_1000": round(gross_lacquer_per_1000, 3),
            "net_total": round(net_lacquer_per_1000 * thousand_units, 2),
            "gross_total": round(gross_lacquer_per_1000 * thousand_units, 2),
            "formula": lacquer_formula
        },
        {
            "item_code": base_item_code,
            "material_group": "External Base Coat",
            "item_name": base_item_name,
            "uom": "kg",
            "net_rate_per_1000": round(net_base_per_1000, 3),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(gross_base_per_1000, 3),
            "net_total": round(net_base_per_1000 * thousand_units, 2),
            "gross_total": round(gross_base_per_1000 * thousand_units, 2),
            "formula": base_formula
        },
        {
            "item_code": varnish_item_code,
            "material_group": "Overprint Varnish",
            "item_name": varnish_item_name,
            "uom": "kg",
            "net_rate_per_1000": round(net_varnish_per_1000, 3),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(gross_varnish_per_1000, 3),
            "net_total": round(net_varnish_per_1000 * thousand_units, 2),
            "gross_total": round(gross_varnish_per_1000 * thousand_units, 2),
            "formula": varnish_formula
        },
        {
            "item_code": "510",
            "material_group": "Printing Ink",
            "item_name": f"SunAltec MB PLUS Dry-Offset Inks ({num_colors} Colors)",
            "uom": "kg",
            "net_rate_per_1000": round(total_net_ink_per_1000, 3),
            "scrap_pct": 10.0,
            "gross_rate_per_1000": round(total_gross_ink_per_1000, 3),
            "net_total": round(total_net_ink_per_1000 * thousand_units, 2),
            "gross_total": round(total_gross_ink_per_1000 * thousand_units, 2),
            "formula": f"{num_colors} colors x 0.280 kg / 1,000 / (1 - 0.10)"
        }
    ]

    estimated_shifts = math.ceil(can_qty / (CANS_PER_SHIFT_NOMINAL * 0.90)) # 90% OEE
    slug_tonnage = (gross_slug_per_1000 * thousand_units) / 1000.0

    return {
        "product_size": product_size,
        "can_qty": can_qty,
        "geometry": size_spec,
        "estimated_shifts": estimated_shifts,
        "slug_tonnage": round(slug_tonnage, 3),
        "items": items,
        "summary": {
            "slug_gross_kg": round(gross_slug_per_1000 * thousand_units, 1),
            "lacquer_gross_kg": round(gross_lacquer_per_1000 * thousand_units, 2),
            "base_coat_gross_kg": round(gross_base_per_1000 * thousand_units, 2),
            "varnish_gross_kg": round(gross_varnish_per_1000 * thousand_units, 2),
            "total_ink_gross_kg": round(total_gross_ink_per_1000 * thousand_units, 2),
            "cleaner_gross_kg": round(gross_cleaner_per_1000 * thousand_units, 1),
        }
    }

def evaluate_shelf_life(received_date_str: str, shelf_life_days: int, as_of_date_str: Optional[str] = None) -> Dict[str, Any]:
    """
    Computes shelf-life countdown and categorization:
    🟢 Normal: < 120 days
    🟡 Warning: 120 <= age < 150 days
    🔴 Critical: >= 150 days (approaching 180d expiration)
    """
    as_of = date.fromisoformat(as_of_date_str) if as_of_date_str else date.today()
    rec = date.fromisoformat(received_date_str)
    age_days = (as_of - rec).days
    remaining_days = max(0, shelf_life_days - age_days)

    if age_days >= 150:
        status = "critical"
        badge_color = "#ef4444" # red
        label = "Critical Shelf-Life (>=150d)"
    elif age_days >= 120:
        status = "warning"
        badge_color = "#f59e0b" # amber
        label = "Warning Shelf-Life (>=120d)"
    else:
        status = "normal"
        badge_color = "#10b981" # emerald
        label = "Normal / Fresh (<120d)"

    return {
        "age_days": age_days,
        "remaining_days": remaining_days,
        "shelf_life_days": shelf_life_days,
        "status": status,
        "badge_color": badge_color,
        "label": label,
        "expiry_date": (rec + timedelta(days=shelf_life_days)).isoformat()
    }

def seed_master_data(db_path: Optional[str] = None, as_of_date: str = "2026-09-12", force_reseed: bool = False, demo_data: bool = False):
    """
    Populates operational records for Alpha Aerosols (Kot Abdul Malik).
    When demo_data=False (default): Sets up the pre-production environment with 'Aerosol Customer',
    0 shift entries, 0 dispatches, and active raw material inventory ready for plant startup.
    When demo_data=True: Seeds historical test shifts and multi-customer orders for automated test suites.
    """
    if db_path is None:
        db_path = DB_PATH
    init_db(db_path)
    conn = get_connection(db_path)
    cur = conn.cursor()

    if force_reseed:
        cur.execute("DELETE FROM dispatches;")
        cur.execute("DELETE FROM shifts;")
        cur.execute("DELETE FROM orders;")
        cur.execute("DELETE FROM inventory;")
        cur.execute("DELETE FROM bom_standards;")
        try:
            cur.execute("DELETE FROM sqlite_sequence WHERE name IN ('orders', 'shifts', 'inventory', 'dispatches', 'bom_standards');")
        except Exception:
            pass
    else:
        # Check if already seeded
        cur.execute("SELECT COUNT(*) FROM orders;")
        if cur.fetchone()[0] > 0:
            conn.close()
            return

    if demo_data:
        # 1. Historical Demo Orders
        orders_data = [
            (
                1,
                "POF-2026-081",
                "Golden Pearl Cosmetics",
                "Golden Pearl Body Spray 150ml",
                "45x160mm",
                120000,
                0.05,
                "2026-08-25",
                "2026-09-18",
                "In Production",
                "GP-BS-150-REV3",
                "Priority export run. High gloss finish required.",
                "2026-08-25T09:00:00"
            ),
            (
                2,
                "POF-2026-082",
                "Alpha Aerosols Stock",
                "Standard Deodorant Can 45x160 White",
                "45x160mm",
                80000,
                0.05,
                "2026-09-01",
                "2026-09-22",
                "In Production",
                "AL-STD-160-W",
                "Buffer stock for spot domestic orders.",
                "2026-09-01T10:30:00"
            ),
            (
                3,
                "POF-2026-083",
                "CareLine Personal Care",
                "CareLine Fresh Deodorant 45x150",
                "45x150mm",
                100000,
                0.05,
                "2026-09-03",
                "2026-09-28",
                "Pending",
                "CL-DEO-150-B",
                "Silk-matt overprint varnish specified. 0.35mm wall.",
                "2026-09-03T14:15:00"
            ),
            (
                4,
                "POF-2026-079",
                "Meditech Aerosols",
                "Meditech Disinfectant Spray 45x160",
                "45x160mm",
                60000,
                0.05,
                "2026-08-10",
                "2026-09-05",
                "Completed",
                "MT-DIS-160-G",
                "Gold protective lacquer double coat specified.",
                "2026-08-10T11:00:00"
            )
        ]

        # 2. Historical Demo Shifts
        shifts_data = [
            ("2026-08-28", "Day", 4, "45x160mm", 14800, 620, 1.0, "Washer nozzle clean", "Tariq Mahmood", "2026-08-28T20:05:00"),
            ("2026-08-28", "Night", 4, "45x160mm", 15100, 580, 0.5, "Ink viscosity adjust", "M. Aslam", "2026-08-29T08:10:00"),
            ("2026-08-29", "Day", 4, "45x160mm", 15250, 510, 0.5, "Minor slug jam at chute", "Tariq Mahmood", "2026-08-29T20:00:00"),
            ("2026-08-29", "Night", 4, "45x160mm", 15400, 460, 0.0, "None", "M. Aslam", "2026-08-30T08:00:00"),
            ("2026-09-08", "Day", 1, "45x160mm", 13800, 780, 1.5, "Print mandrel registration tuning", "Tariq Mahmood", "2026-09-08T20:00:00"),
            ("2026-09-08", "Night", 1, "45x160mm", 14600, 610, 0.75, "OPV coater doctor blade adjustment", "M. Aslam", "2026-09-09T08:00:00"),
            ("2026-09-09", "Day", 1, "45x160mm", 15200, 540, 0.5, "Curing oven temperature stabilization", "Tariq Mahmood", "2026-09-09T20:00:00"),
            ("2026-09-09", "Night", 1, "45x160mm", 15050, 590, 0.5, "Necker guide rail wiper change", "M. Aslam", "2026-09-10T08:00:00"),
            ("2026-09-10", "Day", 1, "45x160mm", 15500, 480, 0.25, "Slug feeder refill", "Tariq Mahmood", "2026-09-10T20:00:00"),
            ("2026-09-10", "Night", 1, "45x160mm", 14900, 520, 0.5, "Printer ink station 2 top-up", "M. Aslam", "2026-09-11T08:00:00"),
            ("2026-09-11", "Day", 1, "45x160mm", 15600, 440, 0.25, "Scheduled line lubrication", "Tariq Mahmood", "2026-09-11T20:00:00"),
            ("2026-09-11", "Night", 1, "45x160mm", 15300, 490, 0.5, "Palletizer sensor cleaning", "M. Aslam", "2026-09-12T08:00:00"),
            ("2026-09-12", "Day", 1, "45x160mm", 8200, 240, 0.25, "Press punch temperature check", "Tariq Mahmood", "2026-09-12T12:30:00"),
        ]

        # 3. Historical Demo Inventory
        inventory_data = [
            ("501", "Aluminum Slugs", "45mm Aluminum Slug (Al99.7%)", "kg", 18500.0, 5000.0, 780.0, "SLUG-2026-07A", "2026-07-15", 1080, "Ambient Dry Store (5°C - 30°C)", "Kot Abdul Malik Central Stores", "2026-09-12T10:00:00"),
            ("502", "Extrusion Lubricant", "SAPILUB LUBRIMET GR8", "kg", 145.0, 30.0, 2400.0, "LUB-2026-03", "2026-06-20", 1080, "Ambient Dry Store (5°C - 30°C)", "Chemical Room Bay 1", "2026-09-12T10:00:00"),
            ("503", "Alkaline Cleaner", "SAPILUB ALULIQUID 13", "kg", 1250.0, 400.0, 850.0, "ALU-2026-02", "2026-06-15", 1080, "Ambient Chemical Bay (5°C - 35°C)", "Washer Chemical Storage", "2026-09-12T10:00:00"),
            ("504", "Internal Lacquer", "SCHEKOSOL INT PROT Gold (400 9 901)", "kg", 860.0, 200.0, 3200.0, "LAC-2026-G1", "2026-07-28", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack A1", "2026-09-12T10:00:00"),
            ("505", "Internal Lacquer", "SCHEKOSOL INT BEIGE (400 4 902)", "kg", 320.0, 150.0, 3400.0, "LAC-2026-B1", "2026-04-12", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack A2", "2026-09-12T10:00:00"),
            ("506", "External Base Coat", "SCHEKOSOL WH BC (422 0 903)", "kg", 1100.0, 300.0, 2100.0, "BC-2026-W2", "2026-08-05", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack B1", "2026-09-12T10:00:00"),
            ("507", "External Base Coat", "SCHEKOSOL CLEAR BC (422 9 918)", "kg", 240.0, 150.0, 2300.0, "BC-2026-C1", "2026-05-10", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack B2", "2026-09-12T10:00:00"),
            ("508", "Overprint Varnish", "SCHEKOSOL OPV GLOSSY (422 9 900)", "kg", 650.0, 200.0, 1900.0, "OPV-2026-G3", "2026-08-12", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack C1", "2026-09-12T10:00:00"),
            ("509", "Overprint Varnish", "SCHEKOSOL OPV SILKMATT (422 9 903)", "kg", 180.0, 150.0, 2250.0, "OPV-2026-S1", "2026-04-18", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack C2", "2026-09-12T10:00:00"),
            ("510", "Printing Ink", "SunAltec MB PLUS 169766 Opaque White", "kg", 85.0, 20.0, 4200.0, "INK-2026-OW", "2026-06-01", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("599", "Printing Ink", "SunAltec MB PLUS 168391 Yellow G/S", "kg", 45.0, 15.0, 4600.0, "INK-2026-Y1", "2026-06-01", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("600", "Printing Ink", "SunAltec MB PLUS 168396 Magenta", "kg", 42.0, 15.0, 4800.0, "INK-2026-M1", "2026-06-01", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("601", "Printing Ink", "SunAltec MB PLUS 168404 Black Conc.", "kg", 55.0, 15.0, 4300.0, "INK-2026-BK", "2026-06-01", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("603", "Printing Ink", "SunAltec MB PLUS 168402 Reflex Blue", "kg", 38.0, 15.0, 4900.0, "INK-2026-RB", "2026-06-01", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
        ]

        # 4. Historical Demo Dispatches
        dispatches_data = [
            ("DC-2026-039", "2026-09-02", 4, 30000, 150, 10, "LES-4412", "Muhammad Akram", "Meditech Aerosols Ltd.", "Delivered", "2026-09-02T16:30:00"),
            ("DC-2026-040", "2026-09-04", 4, 30550, 153, 11, "LZT-8921", "Abdul Rashid", "Meditech Aerosols Ltd.", "Delivered", "2026-09-04T15:45:00"),
            ("DC-2026-041", "2026-09-10", 1, 45000, 225, 15, "LES-4412", "Muhammad Akram", "Golden Pearl Cosmetics", "Delivered", "2026-09-10T17:00:00"),
            ("DC-2026-042", "2026-09-12", 1, 35000, 175, 12, "LZT-8921", "Abdul Rashid", "Golden Pearl Cosmetics", "In Transit", "2026-09-12T11:15:00"),
        ]
    else:
        # Pre-Production Setup: "Aerosol Customer", 0 finished shifts, 0 dispatches
        orders_data = [
            (
                1,
                "POF-2026-001",
                "Aerosol Customer",
                "Aerosol Container 45x160mm",
                "45x160mm",
                100000,
                0.05,
                "2026-09-12",
                "2026-10-15",
                "In Production",
                "AC-CAN-160-REV1",
                "Inaugural pre-production commercial run. Monobloc aluminum container 45x160mm.",
                "2026-09-12T08:00:00"
            )
        ]
        shifts_data = [
            ("2026-09-12", "Day", 1, "45x160mm", 0, 0, 0.0, "Pre-production tooling readiness", "Tariq Mahmood", "2026-09-12T08:00:00")
        ]
        dispatches_data = [
            ("DC-2026-001", "2026-09-12", 1, 0, 0, 0, "Line 1 - Pre-Production", "Pending Assignment", "Aerosol Customer", "Scheduled", "2026-09-12T08:00:00")
        ]
        # Full warehouse inventory ready for plant startup (all fresh batches)
        inventory_data = [
            ("501", "Aluminum Slugs", "45mm Aluminum Slug (Al99.7%)", "kg", 25000.0, 5000.0, 780.0, "SLUG-2026-08A", "2026-08-15", 1080, "Ambient Dry Store (5°C - 30°C)", "Kot Abdul Malik Central Stores", "2026-09-12T10:00:00"),
            ("502", "Extrusion Lubricant", "SAPILUB LUBRIMET GR8", "kg", 200.0, 30.0, 2400.0, "LUB-2026-04", "2026-08-15", 1080, "Ambient Dry Store (5°C - 30°C)", "Chemical Room Bay 1", "2026-09-12T10:00:00"),
            ("503", "Alkaline Cleaner", "SAPILUB ALULIQUID 13", "kg", 1500.0, 400.0, 850.0, "ALU-2026-03", "2026-08-15", 1080, "Ambient Chemical Bay (5°C - 35°C)", "Washer Chemical Storage", "2026-09-12T10:00:00"),
            ("504", "Internal Lacquer", "SCHEKOSOL INT PROT Gold (400 9 901)", "kg", 1000.0, 200.0, 3200.0, "LAC-2026-G2", "2026-08-25", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack A1", "2026-09-12T10:00:00"),
            ("505", "Internal Lacquer", "SCHEKOSOL INT BEIGE (400 4 902)", "kg", 500.0, 150.0, 3400.0, "LAC-2026-B2", "2026-08-25", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack A2", "2026-09-12T10:00:00"),
            ("506", "External Base Coat", "SCHEKOSOL WH BC (422 0 903)", "kg", 1200.0, 300.0, 2100.0, "BC-2026-W3", "2026-08-28", 240, "Air Conditioned (20°C - 25°C)", "Cold Room Rack B1", "2026-09-12T10:00:00"),
            ("507", "External Base Coat", "SCHEKOSOL CLEAR BC (422 9 918)", "kg", 400.0, 150.0, 2300.0, "BC-2026-C2", "2026-08-28", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack B2", "2026-09-12T10:00:00"),
            ("508", "Overprint Varnish", "SCHEKOSOL OPV GLOSSY (422 9 900)", "kg", 800.0, 200.0, 1900.0, "OPV-2026-G4", "2026-09-01", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack C1", "2026-09-12T10:00:00"),
            ("509", "Overprint Varnish", "SCHEKOSOL OPV SILKMATT (422 9 903)", "kg", 350.0, 150.0, 2250.0, "OPV-2026-S2", "2026-09-01", 180, "Air Conditioned (20°C - 25°C)", "Cold Room Rack C2", "2026-09-12T10:00:00"),
            ("510", "Printing Ink", "SunAltec MB PLUS 169766 Opaque White", "kg", 100.0, 20.0, 4200.0, "INK-2026-OW", "2026-08-15", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("599", "Printing Ink", "SunAltec MB PLUS 168391 Yellow G/S", "kg", 50.0, 15.0, 4600.0, "INK-2026-Y1", "2026-08-15", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("600", "Printing Ink", "SunAltec MB PLUS 168396 Magenta", "kg", 50.0, 15.0, 4800.0, "INK-2026-M1", "2026-08-15", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("601", "Printing Ink", "SunAltec MB PLUS 168404 Black Conc.", "kg", 60.0, 15.0, 4300.0, "INK-2026-BK", "2026-08-15", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
            ("603", "Printing Ink", "SunAltec MB PLUS 168402 Reflex Blue", "kg", 50.0, 15.0, 4900.0, "INK-2026-RB", "2026-08-15", 1080, "Ambient Dry Store (15°C - 30°C)", "Ink Dispensing Station", "2026-09-12T10:00:00"),
        ]

    cur.executemany("""
    INSERT INTO orders (
        id, pof_number, customer_name, product_name, product_size, order_qty,
        tolerance_pct, order_date, due_date, status, artwork_ref, notes, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, orders_data)

    if shifts_data:
        cur.executemany("""
        INSERT INTO shifts (
            shift_date, shift_type, pof_id, product_size, good_cans, line_scrap,
            downtime_hours, downtime_reason, supervisor, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, shifts_data)

    cur.executemany("""
    INSERT INTO inventory (
        item_code, category, item_name, uom, balance_qty, min_stock_level,
        unit_cost_pkr, batch_number, received_date, shelf_life_days,
        storage_condition, location, updated_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, inventory_data)

    if dispatches_data:
        cur.executemany("""
        INSERT INTO dispatches (
            challan_number, dispatch_date, pof_id, dispatched_cans,
            carton_count, pallet_count, vehicle_number, driver_name,
            receiver_party, status, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, dispatches_data)

    # 5. BOM Standards
    for size in ["45x160mm", "45x150mm"]:
        bom_calc = calculate_bom(size, 1000)
        for itm in bom_calc["items"]:
            cur.execute("""
            INSERT INTO bom_standards (
                product_size, material_group, item_code, item_name, uom,
                net_rate_per_1000, scrap_rate, gross_rate_per_1000, calculation_rule
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (
                size,
                itm["material_group"],
                itm["item_code"],
                itm["item_name"],
                itm["uom"],
                itm["net_rate_per_1000"],
                itm["scrap_pct"] / 100.0,
                itm["gross_rate_per_1000"],
                itm["formula"]
            ))

    conn.commit()
    conn.close()

def export_production_json(db_path: Optional[str] = None, json_path: Optional[str] = None, as_of_date: Optional[str] = None) -> Dict[str, Any]:
    """
    Generates the standalone JSON snapshot consumed by the frontend.
    Computes real-time KPIs, tolerance bounds, shelf-life indicators,
    and inventory days-of-stock cover.
    """
    if db_path is None:
        db_path = DB_PATH
    if json_path is None:
        json_path = JSON_PATH
    conn = get_connection(db_path)
    cur = conn.cursor()

    ref_date = date.fromisoformat(as_of_date) if as_of_date else date.today()
    ref_date_str = ref_date.isoformat()
    month_prefix = ref_date_str[:7] # e.g. "2026-09"

    # 1. Fetch Orders and calculate production progress
    cur.execute("SELECT * FROM orders ORDER BY id DESC;")
    orders_raw = [dict(r) for r in cur.fetchall()]

    orders = []
    for ord_row in orders_raw:
        cur.execute("""
        SELECT COALESCE(SUM(good_cans), 0) as total_good,
               COALESCE(SUM(line_scrap), 0) as total_scrap
        FROM shifts WHERE pof_id = ?;
        """, (ord_row["id"],))
        shift_stat = cur.fetchone()
        produced_good = shift_stat["total_good"]
        produced_scrap = shift_stat["total_scrap"]
        total_line_run = produced_good + produced_scrap

        cur.execute("""
        SELECT COALESCE(SUM(dispatched_cans), 0) as total_dispatched
        FROM dispatches WHERE pof_id = ?;
        """, (ord_row["id"],))
        disp_stat = cur.fetchone()
        dispatched_total = disp_stat["total_dispatched"]

        order_qty = ord_row["order_qty"]
        tol_pct = ord_row["tolerance_pct"]
        min_qty = math.floor(order_qty * (1.0 - tol_pct))
        max_qty = math.ceil(order_qty * (1.0 + tol_pct))
        completion_pct = round((produced_good / order_qty) * 100.0, 1) if order_qty > 0 else 0.0
        remaining_qty = max(0, order_qty - produced_good)
        scrap_pct = round((produced_scrap / total_line_run) * 100.0, 2) if total_line_run > 0 else 0.0

        orders.append({
            **ord_row,
            "produced_good": produced_good,
            "produced_scrap": produced_scrap,
            "total_line_run": total_line_run,
            "dispatched_total": dispatched_total,
            "fg_stock": max(0, produced_good - dispatched_total),
            "min_acceptable_qty": min_qty,
            "max_acceptable_qty": max_qty,
            "completion_pct": completion_pct,
            "remaining_qty": remaining_qty,
            "order_scrap_pct": scrap_pct,
            "is_within_tolerance": (min_qty <= produced_good <= max_qty) if produced_good >= min_qty else False
        })

    # 2. Fetch Shifts (with joined POF info)
    cur.execute("""
    SELECT s.*, o.pof_number, o.customer_name, o.product_name
    FROM shifts s
    JOIN orders o ON s.pof_id = o.id
    ORDER BY s.shift_date DESC, s.id DESC;
    """)
    shifts_raw = [dict(r) for r in cur.fetchall()]

    shifts = []
    for s in shifts_raw:
        total_cans = s["good_cans"] + s["line_scrap"]
        scrap_pct = round((s["line_scrap"] / total_cans) * 100.0, 2) if total_cans > 0 else 0.0
        shifts.append({
            **s,
            "total_cans": total_cans,
            "scrap_pct": scrap_pct
        })

    # 3. Fetch Inventory and compute live shelf-life and daily burn cover
    cur.execute("SELECT * FROM inventory ORDER BY category ASC, item_code ASC;")
    inventory_raw = [dict(r) for r in cur.fetchall()]

    # Estimate average daily can output from the last 14 recorded production days to compute days-of-stock
    cur.execute("""
    SELECT COALESCE(SUM(good_cans), 0) as total_cans, COUNT(DISTINCT shift_date) as active_days
    FROM shifts
    WHERE shift_date >= date(?, '-14 days');
    """, (ref_date_str,))
    recent = cur.fetchone()
    avg_daily_cans = (recent["total_cans"] / recent["active_days"]) if (recent["active_days"] and recent["active_days"] > 0) else 0.0

    # Rigorous gross consumption reference rates per 1,000 cans for 45x160mm standard container
    ITEM_GROSS_RATES = {
        "501": 22.222,  # 45mm Aluminum Slug (20.0 / 0.90)
        "502": 0.0233,  # SAPILUB LUBRIMET GR8 (0.021 / 0.90)
        "503": 5.556,   # SAPILUB ALULIQUID 13 (5.0 / 0.90)
        "504": 1.6077,  # SCHEKOSOL INT PROT Gold (1.045 / 0.65)
        "505": 1.7538,  # SCHEKOSOL INT BEIGE (1.140 / 0.65)
        "506": 1.0811,  # SCHEKOSOL WH BC (0.973 / 0.90)
        "507": 0.7789,  # SCHEKOSOL CLEAR BC (0.701 / 0.90)
        "508": 0.4778,  # SCHEKOSOL OPV GLOSSY (0.430 / 0.90)
        "509": 0.4278,  # SCHEKOSOL OPV SILKMATT (0.385 / 0.90)
        "510": 0.3111,  # SunAltec Opaque White (0.280 / 0.90 per station)
        "599": 0.3111,  # SunAltec Yellow (0.280 / 0.90 per station)
        "600": 0.3111,  # SunAltec Magenta (0.280 / 0.90 per station)
        "601": 0.3111,  # SunAltec Black (0.280 / 0.90 per station)
        "603": 0.3111   # SunAltec Reflex Blue (0.280 / 0.90 per station)
    }

    inventory = []
    for item in inventory_raw:
        shelf = evaluate_shelf_life(item["received_date"], item["shelf_life_days"], ref_date_str)
        gross_rate = ITEM_GROSS_RATES.get(item["item_code"])
        if gross_rate is None:
            cat = item["category"]
            if "Slug" in cat: gross_rate = 22.222
            elif "Lubricant" in cat: gross_rate = 0.0233
            elif "Cleaner" in cat: gross_rate = 5.556
            elif "Lacquer" in cat: gross_rate = 1.6077
            elif "Base Coat" in cat: gross_rate = 1.0811
            elif "Varnish" in cat: gross_rate = 0.4778
            elif "Ink" in cat: gross_rate = 0.3111
            else: gross_rate = 0.0

        if gross_rate and avg_daily_cans > 0:
            daily_burn = round(gross_rate * (avg_daily_cans / 1000.0), 2)
            days_of_stock = round(item["balance_qty"] / daily_burn, 1) if daily_burn > 0 else 999.0
        else:
            daily_burn = 0.0
            days_of_stock = 999.0

        is_climate_sensitive = "Air Conditioned" in item["storage_condition"]

        inventory.append({
            **item,
            "shelf_evaluation": shelf,
            "daily_burn_rate": daily_burn,
            "days_of_stock": days_of_stock,
            "is_climate_sensitive": is_climate_sensitive,
            "stock_status": "Low Stock" if item["balance_qty"] <= item["min_stock_level"] else "Healthy"
        })

    # 4. Fetch Dispatches
    cur.execute("""
    SELECT d.*, o.pof_number, o.customer_name, o.product_name, o.product_size
    FROM dispatches d
    JOIN orders o ON d.pof_id = o.id
    ORDER BY d.dispatch_date DESC, d.id DESC;
    """)
    dispatches = [dict(r) for r in cur.fetchall()]

    # 5. Master Standards
    standards = {
        "45x160mm": calculate_bom("45x160mm", 100000),
        "45x150mm": calculate_bom("45x150mm", 100000),
    }

    # 6. Aggregated KPIs
    today_shifts = [s for s in shifts if s["shift_date"] == ref_date_str]
    today_good = sum(s["good_cans"] for s in today_shifts)
    today_scrap = sum(s["line_scrap"] for s in today_shifts)
    today_total = today_good + today_scrap
    today_scrap_pct = round((today_scrap / today_total) * 100.0, 2) if today_total > 0 else 0.0

    # Month to Date (MTD)
    mtd_shifts = [s for s in shifts if s["shift_date"].startswith(month_prefix)]
    mtd_good = sum(s["good_cans"] for s in mtd_shifts)
    mtd_scrap = sum(s["line_scrap"] for s in mtd_shifts)
    mtd_total = mtd_good + mtd_scrap
    mtd_scrap_pct = round((mtd_scrap / mtd_total) * 100.0, 2) if mtd_total > 0 else 0.0

    # Today's Dispatches
    today_dispatches = [d for d in dispatches if d["dispatch_date"] == ref_date_str]
    today_dispatched_cans = sum(d["dispatched_cans"] for d in today_dispatches)
    today_dispatches_count = len([d for d in today_dispatches if d.get("dispatched_cans", 0) > 0 or d.get("status") in ("Delivered", "In Transit")])

    # Active POFs
    active_pofs = [o for o in orders if o["status"] in ("In Production", "Pending")]

    # Downtime analysis
    downtime_by_reason = {}
    for s in mtd_shifts:
        reason = s.get("downtime_reason") or "Unspecified"
        if reason and reason != "None" and s.get("downtime_hours", 0.0) > 0:
            downtime_by_reason[reason] = round(downtime_by_reason.get(reason, 0.0) + s["downtime_hours"], 2)

    # Latest Completed Shift (fixing the morning zero trap)
    completed_shifts = [s for s in shifts if (s.get("good_cans", 0) > 0 or s.get("line_scrap", 0) > 0)]
    latest_shift = completed_shifts[0] if completed_shifts else (shifts[0] if shifts else None)

    # FG Warehouse Buffer (Finished Goods awaiting dispatch)
    total_fg_buffer_cans = sum(max(0, o["produced_good"] - o["dispatched_total"]) for o in orders)
    total_fg_buffer_pallets = round(total_fg_buffer_cans / 3000.0, 1) if total_fg_buffer_cans > 0 else 0.0

    # Resolve active monthly workbook
    base_dir = os.path.dirname(os.path.abspath(db_path)) if db_path else os.path.dirname(os.path.abspath(__file__))
    if not os.path.exists(os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")):
        project_root = os.path.dirname(os.path.abspath(__file__))
        if os.path.exists(os.path.join(project_root, "Aerosol_Production_Entry.xlsx")):
            base_dir = project_root
    resolved_wb_path = get_active_aerosol_workbook(base_dir=base_dir, target_date=ref_date)
    active_wb_name = os.path.basename(resolved_wb_path)

    payload = {
        "meta": {
            "plant_name": "Alpha Aerosols",
            "location": "Kot Abdul Malik, Punjab, Pakistan",
            "as_of_date": ref_date_str,
            "active_workbook": active_wb_name,
            "generated_at": datetime.now().isoformat(),
            "version": "2.0.0",
            "sqlite_mode": "WAL"
        },
        "kpis": {
            "today_output_cans": today_good,
            "today_scrap_cans": today_scrap,
            "today_scrap_pct": today_scrap_pct,
            "mtd_output_cans": mtd_good,
            "mtd_scrap_cans": mtd_scrap,
            "mtd_scrap_pct": mtd_scrap_pct,
            "today_dispatches_cans": today_dispatched_cans,
            "today_dispatches_count": today_dispatches_count,
            "active_pofs_count": len(active_pofs),
            "downtime_hours_mtd": round(sum(s["downtime_hours"] for s in mtd_shifts), 1),
            "fg_buffer_cans": total_fg_buffer_cans,
            "fg_buffer_pallets": total_fg_buffer_pallets,
            "latest_shift": {
                "good_cans": latest_shift["good_cans"] if latest_shift else 0,
                "line_scrap": latest_shift["line_scrap"] if latest_shift else 0,
                "scrap_pct": latest_shift["scrap_pct"] if latest_shift else 0.0,
                "shift_date": latest_shift["shift_date"] if latest_shift else ref_date_str,
                "shift_type": latest_shift["shift_type"] if latest_shift else "Day",
                "supervisor": latest_shift["supervisor"] if latest_shift else "Tariq Mahmood",
                "pof_number": latest_shift.get("pof_number") or (active_pofs[0]["pof_number"] if active_pofs else "-"),
                "customer_name": latest_shift.get("customer_name") or (active_pofs[0]["customer_name"] if active_pofs else "-")
            } if latest_shift else None
        },
        "orders": orders,
        "shifts": shifts,
        "inventory": inventory,
        "dispatches": dispatches,
        "standards": standards,
        "downtime_pareto": sorted([{"reason": k, "hours": v} for k, v in downtime_by_reason.items()], key=lambda x: x["hours"], reverse=True)
    }

    conn.close()

    # Write JSON output
    os.makedirs(os.path.dirname(json_path), exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)

    return payload

def add_shift_entry(
    shift_date: str,
    shift_type: str,
    pof_id: int,
    product_size: str,
    good_cans: int,
    line_scrap: int,
    downtime_hours: float,
    downtime_reason: str,
    supervisor: str,
    db_path: Optional[str] = None
) -> int:
    """Inserts a verified single-stage finished shift record and auto-syncs the JSON data layer."""
    if db_path is None:
        db_path = DB_PATH
    conn = get_connection(db_path)
    cur = conn.cursor()
    now_iso = datetime.now().isoformat()

    cur.execute("""
    INSERT INTO shifts (
        shift_date, shift_type, pof_id, product_size, good_cans,
        line_scrap, downtime_hours, downtime_reason, supervisor, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, (
        shift_date, shift_type, pof_id, product_size, good_cans,
        line_scrap, downtime_hours, downtime_reason, supervisor, now_iso
    ))
    shift_id = cur.lastrowid

    # Deduct raw materials consumption from live inventory based on yield-inverse formula
    bom_calc = calculate_bom(product_size, good_cans)
    for itm in bom_calc["items"]:
        consumed_qty = itm["gross_total"]
        cur.execute("""
        UPDATE inventory
        SET balance_qty = MAX(0.0, balance_qty - ?),
            updated_at = ?
        WHERE item_code = ?;
        """, (consumed_qty, now_iso, itm["item_code"]))

    conn.commit()
    conn.close()

    # Automatically refresh JSON export
    export_production_json(db_path)
    return shift_id

def create_order(
    pof_number: str,
    customer_name: str,
    product_name: str,
    product_size: str,
    order_qty: int,
    tolerance_pct: float,
    order_date: str,
    due_date: str,
    artwork_ref: str = "",
    notes: str = "",
    db_path: Optional[str] = None
) -> int:
    """Inserts a new production order."""
    if db_path is None:
        db_path = DB_PATH
    conn = get_connection(db_path)
    cur = conn.cursor()
    now_iso = datetime.now().isoformat()

    cur.execute("""
    INSERT INTO orders (
        pof_number, customer_name, product_name, product_size, order_qty,
        tolerance_pct, order_date, due_date, status, artwork_ref, notes, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'In Production', ?, ?, ?);
    """, (
        pof_number, customer_name, product_name, product_size, order_qty,
        tolerance_pct, order_date, due_date, artwork_ref, notes, now_iso
    ))
    order_id = cur.lastrowid
    conn.commit()
    conn.close()

    export_production_json(db_path)
    return order_id

def create_dispatch(
    challan_number: str,
    dispatch_date: str,
    pof_id: int,
    dispatched_cans: int,
    carton_count: int,
    pallet_count: int,
    vehicle_number: str,
    receiver_party: str,
    driver_name: str = "",
    db_path: Optional[str] = None
) -> int:
    """Inserts a new delivery challan dispatch."""
    if db_path is None:
        db_path = DB_PATH
    conn = get_connection(db_path)
    cur = conn.cursor()
    now_iso = datetime.now().isoformat()

    cur.execute("""
    INSERT INTO dispatches (
        challan_number, dispatch_date, pof_id, dispatched_cans,
        carton_count, pallet_count, vehicle_number, driver_name,
        receiver_party, status, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'Dispatched', ?);
    """, (
        challan_number, dispatch_date, pof_id, dispatched_cans,
        carton_count, pallet_count, vehicle_number, driver_name,
        receiver_party, now_iso
    ))
    dispatch_id = cur.lastrowid
    conn.commit()
    conn.close()

    export_production_json(db_path)
    return dispatch_id

def update_order_status(pof_id: int, status: str, db_path: Optional[str] = None):
    """Updates order status and syncs the JSON data layer."""
    valid_statuses = {"In Production", "Pending", "Completed", "Cancelled", "On Hold"}
    if status not in valid_statuses:
        raise ValueError(f"Invalid status '{status}'. Valid statuses: {valid_statuses}")
    if db_path is None:
        db_path = DB_PATH
    conn = get_connection(db_path)
    try:
        cur = conn.cursor()
        cur.execute("UPDATE orders SET status = ? WHERE id = ?;", (status, pof_id))
        if cur.rowcount == 0:
            raise ValueError(f"Order ID {pof_id} not found.")
        conn.commit()
    finally:
        conn.close()

    export_production_json(db_path)

def export_job_card_excel(pof_id: int, output_path: str, db_path: Optional[str] = None) -> str:
    """
    Generates a shop-floor Job Card Excel workbook cleanly via openpyxl.
    No Excel COM automation, no win32com, zero taskkill.
    """
    if db_path is None:
        db_path = DB_PATH
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    conn = get_connection(db_path)
    try:
        cur = conn.cursor()
        cur.execute("SELECT * FROM orders WHERE id = ?;", (pof_id,))
        row = cur.fetchone()
        if not row:
            raise ValueError(f"Order POF ID {pof_id} not found.")
        order = dict(row)
    finally:
        conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Card"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="12151A", end_color="12151A", fill_type="solid")
    section_fill = PatternFill(start_color="202630", end_color="202630", fill_type="solid")
    th_fill = PatternFill(start_color="2C3440", end_color="2C3440", fill_type="solid")

    font_title = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    font_section = Font(name="Arial", size=11, bold=True, color="F59E0B")
    font_th = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True, color="000000")
    font_regular = Font(name="Arial", size=10, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )

    # Title Block
    ws.merge_cells("A1:F2")
    title_cell = ws["A1"]
    title_cell.value = "ALPHA AEROSOLS — SHOP FLOOR PRODUCTION JOB CARD"
    title_cell.font = font_title
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = "Document #: AER-JC-001"
    ws["A3"].font = font_bold
    ws["D3"] = f"Date Generated: {date.today().isoformat()}"
    ws["D3"].font = font_regular
    ws["F3"] = "Revision: 02"
    ws["F3"].font = font_regular

    # Section 1: Order Details
    ws.merge_cells("A5:F5")
    ws["A5"] = "1. ORDER & PRODUCT SPECIFICATIONS"
    ws["A5"].font = font_section
    ws["A5"].fill = section_fill
    ws["A5"].alignment = Alignment(indent=1)

    order_meta = [
        ("POF Number:", order["pof_number"], "Customer Name:", order["customer_name"]),
        ("Product Name:", order["product_name"], "Container Standard:", order["product_size"]),
        ("Order Quantity:", f"{order['order_qty']:,} cans", "Tolerance Bounds:", f"±{int(order['tolerance_pct']*100)}% ({int(order['order_qty']*(1-order['tolerance_pct'])):,} - {int(order['order_qty']*(1+order['tolerance_pct'])):,})"),
        ("Order Date:", order["order_date"], "Delivery Due Date:", order["due_date"]),
        ("Artwork Reference:", order.get("artwork_ref") or "Approved Art Proof", "Status:", order["status"])
    ]

    r_idx = 6
    for row_data in order_meta:
        ws.cell(row=r_idx, column=1, value=row_data[0]).font = font_bold
        ws.cell(row=r_idx, column=2, value=row_data[1]).font = font_regular
        ws.cell(row=r_idx, column=4, value=row_data[2]).font = font_bold
        ws.cell(row=r_idx, column=5, value=row_data[3]).font = font_regular
        r_idx += 1

    # Section 2: Material Requirements (Physical Yield-Inverse Mass Balance)
    r_idx += 1
    ws.merge_cells(f"A{r_idx}:F{r_idx}")
    ws.cell(row=r_idx, column=1, value="2. REQUIRED RAW MATERIALS (YIELD-INVERSE MASS BALANCE)").font = font_section
    ws.cell(row=r_idx, column=1).fill = section_fill
    ws.cell(row=r_idx, column=1).alignment = Alignment(indent=1)

    r_idx += 1
    headers = ["Item Code", "Material Group", "Item Name & Specification", "Gross Req (kg)", "UOM", "Physical Formula / Transfer Loss"]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=r_idx, column=col_idx, value=h)
        c.font = font_th
        c.fill = th_fill
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5) else "left")

    # Detect specific coating variants from order notes or description
    notes_lower = ((order.get("notes") or "") + " " + (order.get("product_name") or "")).lower()
    bom_options = {}
    if "silk" in notes_lower or "matt" in notes_lower:
        bom_options["varnish_type"] = "silkmatt"
    if "beige" in notes_lower:
        bom_options["lacquer_type"] = "beige"
    if "clear" in notes_lower:
        bom_options["base_coat_type"] = "clear"

    bom = calculate_bom(order["product_size"], order["order_qty"], bom_options)
    for itm in bom["items"]:
        r_idx += 1
        ws.cell(row=r_idx, column=1, value=itm["item_code"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=2, value=itm["material_group"])
        ws.cell(row=r_idx, column=3, value=itm["item_name"])
        ws.cell(row=r_idx, column=4, value=itm["gross_total"]).alignment = Alignment(horizontal="right")
        ws.cell(row=r_idx, column=5, value=itm["uom"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=6, value=itm["formula"])

        for c_i in range(1, 7):
            ws.cell(row=r_idx, column=c_i).border = thin_border
            ws.cell(row=r_idx, column=c_i).font = font_regular

    # Section 3: Signatures & Approvals
    r_idx += 2
    ws.merge_cells(f"A{r_idx}:F{r_idx}")
    ws.cell(row=r_idx, column=1, value="3. SHOP FLOOR QUALITY & ISSUANCE SIGN-OFF").font = font_section
    ws.cell(row=r_idx, column=1).fill = section_fill

    r_idx += 2
    ws.cell(row=r_idx, column=1, value="Prepared By (Planning):").font = font_bold
    ws.cell(row=r_idx, column=2, value="Sikander (Planning Lead)").font = font_regular
    ws.cell(row=r_idx, column=4, value="Approved By:").font = font_bold
    ws.cell(row=r_idx, column=5, value="Production Manager").font = font_regular

    r_idx += 2
    ws.cell(row=r_idx, column=1, value="Material Issued (Stores):").font = font_bold
    ws.cell(row=r_idx, column=2, value="__________________").font = font_regular
    ws.cell(row=r_idx, column=4, value="Floor Line Lead:").font = font_bold
    ws.cell(row=r_idx, column=5, value="__________________").font = font_regular

    # Column Widths
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 45

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path

def parse_clean_number(val: Any) -> Optional[float]:
    """Parses numeric string or float safely, handling commas, percentages, k/m multipliers, and Excel error strings."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if math.isnan(val):
            return None
        return float(val)
    val_str = str(val).strip().strip('"\'')
    if not val_str or val_str.lower() in ("none", "null", "nan", "-", "", "#value!", "#ref!", "#n/a"):
        return None
    cleaned = val_str.replace(",", "").replace(" ", "").replace("pcs", "").strip()
    if cleaned.endswith("%"):
        try:
            return float(cleaned[:-1].strip())
        except ValueError:
            return None
    if cleaned.lower().endswith("k"):
        try:
            return float(cleaned[:-1].strip()) * 1000.0
        except ValueError:
            return None
    if cleaned.lower().endswith("m"):
        try:
            return float(cleaned[:-1].strip()) * 1000000.0
        except ValueError:
            return None
    try:
        return float(cleaned)
    except ValueError:
        return None

def parse_clean_date(val: Any) -> Optional[str]:
    """Parses various date formats, timestamps, Excel date serials, and date objects into standard ISO-8601 YYYY-MM-DD string."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)):
        # Excel date serial (days since 1899-12-30)
        if 20000 <= val <= 80000:
            try:
                dt_serial = date(1899, 12, 30) + timedelta(days=int(val))
                return dt_serial.strftime("%Y-%m-%d")
            except Exception:
                pass
        return None
    val_str = str(val).strip().strip('"\'')
    if not val_str:
        return None
    # Fast match ISO YYYY-MM-DD
    m_iso = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})", val_str)
    if m_iso:
        try:
            return date(int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))).strftime("%Y-%m-%d")
        except ValueError:
            pass
    for fmt in (
        "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
        "%d/%m/%Y", "%d/%m/%Y %H:%M:%S",
        "%d-%m-%Y", "%d-%m-%Y %H:%M:%S",
        "%Y/%m/%d", "%Y/%m/%d %H:%M:%S",
        "%d-%b-%Y", "%d-%B-%Y", "%d/%b/%Y", "%d/%B/%Y",
        "%d.%m.%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y",
        "%b %d, %Y", "%B %d, %Y", "%Y.%m.%d"
    ):
        try:
            return datetime.strptime(val_str[:19], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None

def normalize_pof(val: Any) -> str:
    """
    Normalizes various POF representations into a canonical comparable key.
    E.g.: 'POF-2026-081', 'POF-081', '81', 81, 'POF 81', 'POF # 81' -> '81'
    'POF-2026-001' -> '1'
    'POF-SPECIAL-A' -> 'SPECIAL-A'
    """
    if val is None:
        return ""
    s = str(val).strip().upper()
    if not s:
        return ""
    s = re.sub(r"^POF\s*#?[\s\-_]*", "", s)
    m = re.match(r"^\d{4}[\-_](\d+)$", s)
    if m:
        return str(int(m.group(1)))
    if s.isdigit():
        return str(int(s))
    return s

def get_monthly_workbook_name(target_date: Optional[Union[str, date, datetime]] = None) -> str:
    """
    Returns monthly workbook name formatted as Aerosol_MmmYY.xlsx (e.g., Aerosol_Sep26.xlsx).
    Matches the Tubex standard monthly convention (Tubex_MmmYY.xlsx).
    """
    if target_date is None:
        dt = datetime.now()
    elif isinstance(target_date, datetime):
        dt = target_date
    elif isinstance(target_date, date):
        dt = datetime.combine(target_date, datetime.min.time())
    elif isinstance(target_date, str):
        target_str = target_date.strip().strip('"\'')
        dt = None
        for fmt in (
            "%B %Y", "%b %Y", "%B %y", "%b %y", "%B%Y", "%b%y", "%b-%y", "%B-%y", "%B-%Y",
            "%Y-%m-%d", "%Y-%m", "%d-%b-%Y", "%d/%m/%Y"
        ):
            try:
                dt = datetime.strptime(target_str, fmt)
                break
            except Exception:
                if len(target_str) >= 10:
                    try:
                        dt = datetime.strptime(target_str[:10], fmt)
                        break
                    except Exception:
                        pass
        if dt is None:
            m = re.search(r"([A-Za-z]{3,9})[\s\-_]*(\d{2,4})", target_str)
            if m:
                for m_fmt in ("%B", "%b"):
                    for y_fmt in ("%Y", "%y"):
                        try:
                            dt = datetime.strptime(f"{m.group(1)} {m.group(2)}", f"{m_fmt} {y_fmt}")
                            break
                        except Exception:
                            pass
                    if dt:
                        break
        if dt is None:
            dt = datetime.now()
    else:
        dt = datetime.now()

    return f"Aerosol_{dt.strftime('%b%y')}.xlsx"

def get_active_aerosol_workbook(
    base_dir: Optional[str] = None,
    target_date: Optional[Union[str, date, datetime]] = None,
    create_if_missing: bool = False
) -> str:
    r"""
    Resolves the active monthly Aerosol workbook following the Tubex convention.
    Priority hierarchy:
    1. Exact target month file (e.g. Aerosol_Sep26.xlsx) in base_dir.
    2. If create_if_missing is True, initializes target month workbook from template or blank.
    3. Chronologically latest monthly file matching Aerosol_[A-Za-z]{3}\d{2}.xlsx (excluding temp lock ~$ files).
    4. Fallback master template: Aerosol_Production_Entry.xlsx in base_dir.
    """
    if base_dir is None:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        # Ensure base_dir points to project root where Excel files reside
        if not os.path.exists(os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")):
            project_root = os.path.dirname(os.path.abspath(__file__))
            if os.path.exists(os.path.join(project_root, "Aerosol_Production_Entry.xlsx")):
                base_dir = project_root

    target_name = get_monthly_workbook_name(target_date)
    exact_path = os.path.join(base_dir, target_name)

    # 1. Exact monthly file exists
    if os.path.exists(exact_path):
        return exact_path

    # 2. If create_if_missing is True, create exact_path from template
    if create_if_missing:
        fallback_tmpl = os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")
        if os.path.exists(fallback_tmpl):
            try:
                shutil.copyfile(fallback_tmpl, exact_path)
                return exact_path
            except Exception:
                pass
        # If no template exists, create blank workbook with Data_Entry
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data_Entry"
        ws.cell(row=1, column=1, value="AEROSOL PLANT — DAILY PRODUCTION DATA ENTRY")
        headers = [
            'Date', 'Machine', 'POF #', 'Product Name', 'PID', 'Customer',
            'Total Production\n(pcs)', 'Good Production\n(pcs)', 'Rejects\n(pcs)',
            'Rejection\n%', 'DownTime', 'Remarks'
        ]
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)
        wb.save(exact_path)
        return exact_path

    # 3. Check for existing monthly workbooks (sorted chronologically)
    pattern = re.compile(r"^Aerosol_([A-Za-z]{3}\d{2})\.xlsx$", re.IGNORECASE)
    candidates = []
    for f in glob.glob(os.path.join(base_dir, "Aerosol_*.xlsx")):
        fn = os.path.basename(f)
        if fn.startswith("~$"):
            continue
        m = pattern.match(fn)
        if m:
            try:
                dt_cand = datetime.strptime(m.group(1), "%b%y")
            except Exception:
                dt_cand = datetime.min
            candidates.append((dt_cand, f))

    if candidates:
        candidates.sort(key=lambda x: x[0])
        return candidates[-1][1]

    # 4. Fallback to Aerosol_Production_Entry.xlsx
    fallback_path = os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")
    if os.path.exists(fallback_path):
        return fallback_path

    return exact_path

def sync_excel_to_db(
    excel_path: Optional[str] = None,
    db_path: Optional[str] = None,
    json_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Option B: Synchronizes shift production records from an Excel workbook into aerosol.db (SQLite WAL).
    Enforces Zero Windows COM Automation (pure Python openpyxl).
    Deducts raw materials via physical yield-inverse mass balance.
    Auto-refreshes production JSON snapshot (respects custom json_path).
    Supports dynamic monthly files (Aerosol_MmmYY.xlsx) and fallback (Aerosol_Production_Entry.xlsx).
    """
    if db_path is None:
        db_path = DB_PATH

    base_dir = os.path.dirname(os.path.abspath(db_path)) if db_path else os.path.dirname(os.path.abspath(__file__))
    if not os.path.exists(os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")):
        project_root = os.path.dirname(os.path.abspath(__file__))
        if os.path.exists(os.path.join(project_root, "Aerosol_Production_Entry.xlsx")):
            base_dir = project_root

    if excel_path is None:
        excel_path = get_active_aerosol_workbook(base_dir=base_dir)
    elif not os.path.isabs(excel_path):
        excel_path = os.path.join(base_dir, excel_path)

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel workbook not found: {excel_path}")

    import openpyxl

    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Choose target sheet
    target_sheet_name = None
    for name in ["Data_Entry", "Production_Log", "Daily_Production", "Production", "Shifts"]:
        if name in wb.sheetnames:
            target_sheet_name = name
            break
    ws = wb[target_sheet_name] if target_sheet_name else wb.active

    # Identify header row (scanning first 8 rows)
    header_row_idx = None
    headers = {}
    for r in range(1, min(9, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(25, ws.max_column + 1))]
        str_vals = [str(v).lower().strip() for v in row_vals if v is not None]
        if any("date" in s for s in str_vals) and (any("good" in s or "total" in s or "pof" in s or "machine" in s or "product" in s or "target" in s for s in str_vals)):
            header_row_idx = r
            for c in range(1, min(25, ws.max_column + 1)):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    clean_header = str(val).lower().replace('\n', ' ').strip()
                    headers[clean_header] = c
            break

    if not header_row_idx:
        return {
            "status": "error",
            "message": f"Could not identify valid header row in Excel worksheet ({excel_path}).",
            "shifts_imported": 0,
            "skipped_duplicates": 0,
            "skipped_empty": 0
        }

    def get_col_val(row_idx, keywords, exclude_keywords=None):
        for h_text, col_idx in headers.items():
            if any(k in h_text for k in keywords):
                if exclude_keywords and any(ek in h_text for ek in exclude_keywords):
                    continue
                return ws.cell(row=row_idx, column=col_idx).value
        return None

    conn = get_connection(db_path)
    cur = conn.cursor()

    shifts_imported = 0
    shifts_updated = 0
    skipped_duplicates = 0
    skipped_empty = 0

    now_iso = datetime.now().isoformat()

    for r in range(header_row_idx + 1, ws.max_row + 1):
        raw_date = get_col_val(r, ["date"])
        raw_pof = get_col_val(r, ["pof", "order"])
        raw_product = get_col_val(r, ["product"])
        raw_pid = get_col_val(r, ["pid", "product id"])
        raw_machine = get_col_val(r, ["machine", "line"])
        raw_customer = get_col_val(r, ["customer", "client"])
        raw_good = get_col_val(r, ["good"], exclude_keywords=["%", "pct", "rate"])
        raw_rejects = get_col_val(r, ["reject", "scrap", "waste"], exclude_keywords=["%", "pct", "rate"])
        raw_reject_pct = get_col_val(r, ["rejection %", "reject %", "scrap %", "waste%"])
        raw_total = get_col_val(r, ["total", "target"], exclude_keywords=["%", "pct", "rate", "down"])
        raw_remarks = get_col_val(r, ["remark", "reason", "cause", "notes"])
        raw_supervisor = get_col_val(r, ["supervisor", "lead", "operator"])
        raw_shift = get_col_val(r, ["shift"])

        # Downtime parsing: match downtime / dt columns with word boundaries (avoid false positives like 'width')
        total_dt = 0.0
        dt_subtypes = []
        for h_text, col_idx in headers.items():
            if re.search(r"\b(dt|down\s*time)\b", h_text, re.IGNORECASE) and not re.search(r"\b(reason|remark|cause|notes?)\b", h_text, re.IGNORECASE):
                cell_v = ws.cell(row=r, column=col_idx).value
                parsed_v = parse_clean_number(cell_v)
                if parsed_v and parsed_v > 0:
                    is_mins = bool(re.search(r"\b(min|mins|minutes?)\b", h_text, re.IGNORECASE))
                    dt_val_hrs = round(parsed_v / 60.0, 2) if (is_mins or parsed_v > 12.0) else round(parsed_v, 2)
                    total_dt += dt_val_hrs
                    col_title = re.sub(r"\bdt\b", "DT", h_text, flags=re.IGNORECASE).title()
                    dt_subtypes.append(f"{col_title}: {dt_val_hrs}h")

        # Check if row is empty or template row
        if raw_date is None and raw_good is None and raw_total is None and raw_pof is None:
            skipped_empty += 1
            continue

        # Parse date
        shift_date = parse_clean_date(raw_date)
        if not shift_date:
            skipped_empty += 1
            continue

        # Parse and reconcile numbers
        good_num = parse_clean_number(raw_good)
        scrap_num = parse_clean_number(raw_rejects)
        total_num = parse_clean_number(raw_total)

        # If scrap was not in count column, check if scrap rate exists
        if scrap_num is None and raw_reject_pct is not None:
            pct_val = parse_clean_number(raw_reject_pct)
            if pct_val is not None and total_num is not None:
                rate = pct_val / 100.0 if pct_val > 1.0 else pct_val
                scrap_num = round(total_num * rate)

        if good_num is not None and scrap_num is not None:
            good_cans = int(round(good_num))
            line_scrap = int(round(scrap_num))
        elif good_num is not None and total_num is not None:
            good_cans = int(round(good_num))
            line_scrap = max(0, int(round(total_num - good_num)))
        elif scrap_num is not None and total_num is not None:
            line_scrap = int(round(scrap_num))
            good_cans = max(0, int(round(total_num - scrap_num)))
        elif good_num is not None:
            good_cans = int(round(good_num))
            line_scrap = 0
        elif total_num is not None:
            good_cans = int(round(total_num))
            line_scrap = 0
        else:
            good_cans = 0
            line_scrap = 0

        pof_str = str(raw_pof).strip() if raw_pof is not None else ""

        # Only skip if 0 production AND no downtime AND no POF was specified
        if good_cans == 0 and line_scrap == 0 and total_dt <= 0 and not pof_str:
            skipped_empty += 1
            continue

        # Determine product size from PID, product name, or remarks
        combined_meta = f"{raw_pid or ''} {raw_product or ''} {raw_remarks or ''}"
        if "9003" in str(raw_pid or "") or "5003" in str(raw_pid or "") or "150" in combined_meta:
            inferred_size = "45x150mm" if "150" in combined_meta else "45x160mm"
        else:
            inferred_size = "45x160mm"

        # Resolve POF and Order using canonical normalization
        pof_id = None
        product_size = None
        norm_pof_key = normalize_pof(pof_str)

        if pof_str or norm_pof_key:
            cur.execute("SELECT id, pof_number, customer_name, product_size FROM orders ORDER BY id ASC;")
            all_orders = [dict(row) for row in cur.fetchall()]
            for ord_row in all_orders:
                if (norm_pof_key and normalize_pof(ord_row["pof_number"]) == norm_pof_key) or str(ord_row["id"]) == str(pof_str) or ord_row["pof_number"].upper() == pof_str.upper():
                    pof_id = ord_row["id"]
                    product_size = ord_row["product_size"]
                    break

        cust_str = str(raw_customer).strip() if raw_customer else ""
        if not pof_id and cust_str:
            cur.execute("SELECT id, product_size FROM orders WHERE LOWER(customer_name) LIKE LOWER(?) ORDER BY id DESC;", (f"%{cust_str}%",))
            matched_order = cur.fetchone()
            if matched_order:
                pof_id = matched_order["id"]
                product_size = matched_order["product_size"]

        if not pof_id:
            # If pof_str was provided but didn't exist in DB, create this new order!
            if pof_str:
                clean_pof = pof_str if pof_str.upper().startswith("POF-") else f"POF-{pof_str}"
                product_size = inferred_size
                p_name = str(raw_product).strip() if raw_product else f"Aerosol Container {product_size}"
                cur.execute("""
                INSERT INTO orders (pof_number, customer_name, product_name, product_size, order_qty, tolerance_pct, order_date, due_date, status, created_at)
                VALUES (?, ?, ?, ?, 100000, 0.05, ?, ?, 'In Production', ?);
                """, (clean_pof, cust_str or "Aerosol Customer", p_name, product_size, shift_date, shift_date, now_iso))
                pof_id = cur.lastrowid
            else:
                # Check for active order with matching product size
                cur.execute("SELECT id, product_size FROM orders WHERE status IN ('In Production', 'Pending') AND product_size = ? ORDER BY id ASC LIMIT 1;", (inferred_size,))
                active_ord = cur.fetchone()
                if active_ord:
                    pof_id = active_ord["id"]
                    product_size = active_ord["product_size"]
                else:
                    cur.execute("SELECT id, product_size FROM orders ORDER BY id ASC LIMIT 1;")
                    first_ord = cur.fetchone()
                    if first_ord:
                        pof_id = first_ord["id"]
                        product_size = first_ord["product_size"]
                    else:
                        product_size = inferred_size
                        cur.execute("""
                        INSERT INTO orders (pof_number, customer_name, product_name, product_size, order_qty, tolerance_pct, order_date, due_date, status, created_at)
                        VALUES ('POF-2026-001', 'Aerosol Customer', 'Aerosol Container 45x160mm', ?, 100000, 0.05, ?, ?, 'In Production', ?);
                        """, (product_size, shift_date, shift_date, now_iso))
                        pof_id = cur.lastrowid

        if not product_size:
            product_size = inferred_size

        # Shift type (A/1/Day -> Day; B/2/Night -> Night)
        shift_str = str(raw_shift).strip().lower() if raw_shift is not None else ""
        remarks_str = str(raw_remarks).lower() if raw_remarks else ""
        if shift_str in ("b", "night", "n", "2", "shift b", "shift 2") or "night" in remarks_str or "shift b" in remarks_str:
            shift_type = "Night"
        else:
            shift_type = "Day"

        # Downtime hours (capped at 12.0 hours for single shift duration)
        downtime_hours = min(12.0, max(0.0, round(total_dt, 2)))

        # Supervisor resolution
        supervisor = "Tariq Mahmood"
        if raw_supervisor:
            supervisor = str(raw_supervisor).strip()
        elif raw_remarks:
            for s_name in ["Tariq Mahmood", "M. Aslam", "Sikander", "Production Manager"]:
                if s_name.lower() in str(raw_remarks).lower():
                    supervisor = s_name
                    break

        # Clean downtime reason
        downtime_reason = "None"
        if raw_remarks:
            rem_clean = str(raw_remarks).strip()
            if rem_clean.lower() == supervisor.lower():
                downtime_reason = "Normal continuous run" if downtime_hours == 0 else "Unspecified stoppage"
            else:
                for s_name in ["Tariq Mahmood", "M. Aslam", "Sikander", "Production Manager"]:
                    rem_clean = rem_clean.replace(f"- {s_name}", "").replace(f"-{s_name}", "").strip()
                downtime_reason = rem_clean or ("Normal continuous run" if downtime_hours == 0 else "Unspecified stoppage")
        elif dt_subtypes:
            downtime_reason = "; ".join(dt_subtypes)
        elif raw_machine and downtime_hours > 0:
            downtime_reason = f"{raw_machine} stoppage"

        # Check if shift already exists for this (date, shift_type, pof_id)
        cur.execute("""
        SELECT id, good_cans, line_scrap, downtime_hours, downtime_reason, supervisor, product_size
        FROM shifts
        WHERE shift_date = ? AND shift_type = ? AND pof_id = ?;
        """, (shift_date, shift_type, pof_id))
        existing_shift = cur.fetchone()

        if existing_shift:
            old_good = existing_shift["good_cans"]
            old_scrap = existing_shift["line_scrap"]
            old_dt = existing_shift["downtime_hours"]

            # If identical within tolerance, skip
            if old_good == good_cans and old_scrap == line_scrap and abs(old_dt - downtime_hours) < 0.01:
                skipped_duplicates += 1
                continue

            # Existing shift has updated quantities or downtime in Excel!
            diff_good = good_cans - old_good
            cur.execute("""
            UPDATE shifts
            SET good_cans = ?, line_scrap = ?, downtime_hours = ?, downtime_reason = ?, supervisor = ?, product_size = ?
            WHERE id = ?;
            """, (good_cans, line_scrap, downtime_hours, downtime_reason, supervisor, product_size, existing_shift["id"]))

            # Reconcile raw material consumption delta via physical mass balance
            if diff_good != 0:
                delta_bom = calculate_bom(product_size, abs(diff_good))
                for itm in delta_bom["items"]:
                    delta_qty = itm["gross_total"]
                    if diff_good > 0:
                        cur.execute("""
                        UPDATE inventory
                        SET balance_qty = MAX(0.0, balance_qty - ?), updated_at = ?
                        WHERE item_code = ?;
                        """, (delta_qty, now_iso, itm["item_code"]))
                    else:
                        cur.execute("""
                        UPDATE inventory
                        SET balance_qty = balance_qty + ?, updated_at = ?
                        WHERE item_code = ?;
                        """, (delta_qty, now_iso, itm["item_code"]))

            shifts_updated += 1
        else:
            # Insert new shift
            cur.execute("""
            INSERT INTO shifts (
                shift_date, shift_type, pof_id, product_size, good_cans,
                line_scrap, downtime_hours, downtime_reason, supervisor, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (shift_date, shift_type, pof_id, product_size, good_cans, line_scrap, downtime_hours, downtime_reason, supervisor, now_iso))

            # Deduct raw material consumption via yield-inverse mass balance
            bom_calc = calculate_bom(product_size, good_cans)
            for itm in bom_calc["items"]:
                consumed_qty = itm["gross_total"]
                cur.execute("""
                UPDATE inventory
                SET balance_qty = MAX(0.0, balance_qty - ?),
                    updated_at = ?
                WHERE item_code = ?;
                """, (consumed_qty, now_iso, itm["item_code"]))

            shifts_imported += 1

    conn.commit()
    conn.close()

    # Re-export JSON snapshot (respecting custom json_path if provided)
    export_production_json(db_path, json_path=json_path)

    return {
        "status": "success",
        "file": excel_path,
        "filename": os.path.basename(excel_path),
        "shifts_imported": shifts_imported,
        "shifts_updated": shifts_updated,
        "skipped_duplicates": skipped_duplicates,
        "skipped_empty": skipped_empty
    }

def sync_db_to_excel(
    excel_path: Optional[str] = None,
    db_path: Optional[str] = None,
    month_str: Optional[str] = None,
    filter_date: Optional[str] = None
) -> Dict[str, Any]:
    """
    Option A: Synchronizes production shifts from aerosol.db (SQLite) into an Excel workbook.
    Enforces Zero Windows COM Automation (pure openpyxl).
    Preserves existing formatting, sheet formulas, and structures.
    Idempotent: Appends only new unrecorded shifts into the first available rows.
    """
    if db_path is None:
        db_path = DB_PATH

    base_dir = os.path.dirname(os.path.abspath(db_path)) if db_path else os.path.dirname(os.path.abspath(__file__))
    if not os.path.exists(os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")):
        project_root = os.path.dirname(os.path.abspath(__file__))
        if os.path.exists(os.path.join(project_root, "Aerosol_Production_Entry.xlsx")):
            base_dir = project_root

    if excel_path is None:
        if month_str:
            excel_path = os.path.join(base_dir, f"Aerosol_{month_str}.xlsx")
            if not os.path.exists(excel_path):
                tmpl = os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")
                if os.path.exists(tmpl):
                    shutil.copyfile(tmpl, excel_path)
                else:
                    get_active_aerosol_workbook(base_dir=base_dir, create_if_missing=True)
        else:
            excel_path = get_active_aerosol_workbook(base_dir=base_dir, create_if_missing=True)
    elif not os.path.isabs(excel_path):
        excel_path = os.path.join(base_dir, excel_path)

    if not os.path.exists(excel_path):
        tmpl = os.path.join(base_dir, "Aerosol_Production_Entry.xlsx")
        if os.path.exists(tmpl):
            shutil.copyfile(tmpl, excel_path)
        else:
            get_active_aerosol_workbook(base_dir=base_dir, create_if_missing=True)

    import openpyxl

    wb = openpyxl.load_workbook(excel_path)

    target_sheet_name = None
    for name in ["Data_Entry", "Production_Log", "Daily_Production", "Production", "Shifts"]:
        if name in wb.sheetnames:
            target_sheet_name = name
            break
    ws = wb[target_sheet_name] if target_sheet_name else wb.active

    # Identify header row
    header_row_idx = None
    headers = {}
    for r in range(1, min(9, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(25, ws.max_column + 1))]
        str_vals = [str(v).lower().strip() for v in row_vals if v is not None]
        if any("date" in s for s in str_vals) and (any("good" in s or "total" in s or "pof" in s or "machine" in s or "product" in s for s in str_vals)):
            header_row_idx = r
            for c in range(1, min(25, ws.max_column + 1)):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    clean_header = str(val).lower().replace('\n', ' ').strip()
                    headers[clean_header] = c
            break

    # If no header row found, establish default standard schema on row 2
    if not header_row_idx:
        header_row_idx = 2
        ws.cell(row=1, column=1, value="AEROSOL PLANT — DAILY PRODUCTION DATA ENTRY")
        default_headers = [
            'Date', 'Machine', 'POF #', 'Product Name', 'PID', 'Customer',
            'Total Production\n(pcs)', 'Good Production\n(pcs)', 'Rejects\n(pcs)',
            'Rejection\n%', 'DownTime', 'Remarks'
        ]
        for c_idx, h_text in enumerate(default_headers, 1):
            ws.cell(row=header_row_idx, column=c_idx, value=h_text)
            clean_header = h_text.lower().replace('\n', ' ').strip()
            headers[clean_header] = c_idx

    def get_target_col(keywords, exclude_keywords=None, default_col=1):
        for h_text, col_idx in headers.items():
            if any(k in h_text for k in keywords):
                if exclude_keywords and any(ek in h_text for ek in exclude_keywords):
                    continue
                return col_idx
        return default_col

    col_date = get_target_col(["date"], default_col=1)
    col_shift = get_target_col(["shift"], default_col=None)
    col_machine = get_target_col(["machine", "line"], default_col=2)
    col_pof = get_target_col(["pof", "order"], default_col=3)
    col_product = get_target_col(["product"], default_col=4)
    col_pid = get_target_col(["pid", "product id"], default_col=5)
    col_customer = get_target_col(["customer", "client"], default_col=6)
    col_total = get_target_col(["total", "target"], exclude_keywords=["%", "pct", "rate", "down"], default_col=7)
    col_good = get_target_col(["good"], exclude_keywords=["%", "pct", "rate"], default_col=8)
    col_scrap = get_target_col(["reject", "scrap", "waste"], exclude_keywords=["%", "pct", "rate"], default_col=9)
    col_scrap_pct = get_target_col(["rejection %", "reject %", "scrap %", "waste%"], default_col=10)
    col_downtime = get_target_col(["down", "dt"], exclude_keywords=["reason", "remark", "cause"], default_col=11)
    col_remarks = get_target_col(["remark", "reason", "cause", "notes"], default_col=12)

    # Scan existing rows for recorded shifts and find last non-empty row
    existing_by_pof = {}
    existing_by_date_shift = {}
    last_data_row = header_row_idx

    for r in range(header_row_idx + 1, ws.max_row + 1):
        d_val = ws.cell(row=r, column=col_date).value
        g_val = ws.cell(row=r, column=col_good).value
        p_val = ws.cell(row=r, column=col_pof).value
        s_val = ws.cell(row=r, column=col_scrap).value
        rem_val = ws.cell(row=r, column=col_remarks).value
        tot_val = ws.cell(row=r, column=col_total).value
        dt_val = ws.cell(row=r, column=col_downtime).value

        # Row is empty if Date, Good, Total, and POF are all blank
        if (d_val is None or str(d_val).strip() == "") and \
           (g_val is None or str(g_val).strip() == "") and \
           (tot_val is None or str(tot_val).strip() == "") and \
           (p_val is None or str(p_val).strip() == ""):
            continue

        last_data_row = max(last_data_row, r)

        clean_d = parse_clean_date(d_val)
        if clean_d:
            clean_g = int(round(parse_clean_number(g_val) or 0))
            clean_s = int(round(parse_clean_number(s_val) or 0))
            clean_dt = float(round(parse_clean_number(dt_val) or 0.0, 2))
            norm_p = normalize_pof(p_val)
            if col_shift:
                sh_val = ws.cell(row=r, column=col_shift).value
                sh_str = str(sh_val or "").lower().strip()
                st = "Night" if sh_str in ("night", "b", "2", "shift b", "shift 2") or "night" in sh_str else "Day"
            else:
                rem_str = str(rem_val or "").lower()
                st = "Night" if "night" in rem_str else "Day"

            entry = {"row": r, "good": clean_g, "scrap": clean_s, "dt": clean_dt, "norm_pof": norm_p}
            if norm_p:
                existing_by_pof[(clean_d, st, norm_p)] = entry
            if (clean_d, st) not in existing_by_date_shift:
                existing_by_date_shift[(clean_d, st)] = []
            existing_by_date_shift[(clean_d, st)].append(entry)

    # Read shifts from SQLite
    conn = get_connection(db_path)
    cur = conn.cursor()
    cur.execute("""
    SELECT s.id, s.shift_date, s.shift_type, s.pof_id, s.product_size,
           s.good_cans, s.line_scrap, s.downtime_hours, s.downtime_reason,
           s.supervisor, s.created_at,
           o.pof_number, o.customer_name, o.product_name
    FROM shifts s
    LEFT JOIN orders o ON s.pof_id = o.id
    ORDER BY s.shift_date ASC, s.id ASC;
    """)
    shifts = [dict(row) for row in cur.fetchall()]
    conn.close()

    # Determine monthly filter if target filename has monthly code (e.g. Aerosol_Sep26.xlsx -> 2026-09)
    month_filter = None
    fn = os.path.basename(excel_path)
    m_code = re.search(r"Aerosol_([A-Za-z]{3}\d{2})\.xlsx", fn, re.IGNORECASE)
    if m_code:
        try:
            mmm_yy = m_code.group(1)
            dt_m = datetime.strptime(mmm_yy, "%b%y")
            month_filter = dt_m.strftime("%Y-%m")
        except Exception:
            pass

    rows_written = 0
    rows_updated = 0

    for s in shifts:
        s_date = s["shift_date"]
        if filter_date and s_date != filter_date:
            continue
        if month_filter and not s_date.startswith(month_filter):
            continue

        s_good = s["good_cans"] or 0
        s_scrap = s["line_scrap"] or 0
        s_total = s_good + s_scrap
        s_pof = s["pof_number"] or f"POF-{s['pof_id']}"
        s_norm_pof = normalize_pof(s_pof)
        s_st = s["shift_type"] or "Day"
        s_dt = round(float(s["downtime_hours"] or 0.0), 2)

        prod_size = s["product_size"] or "45x160mm"
        pid = 5002 if "160" in prod_size else 5003
        p_name = s["product_name"] or f"AEROSOL CAN {prod_size} PRINTED"
        if not str(p_name).startswith(str(pid)):
            display_prod = f"{pid} - {p_name}"
        else:
            display_prod = p_name

        scrap_rate = round(s_scrap / s_total, 4) if s_total > 0 else 0.0
        rem = f"{s_st} Shift - {s['supervisor'] or 'Line Lead'}"
        dt_reason = s.get("downtime_reason")
        if dt_reason and dt_reason.lower() not in ("none", "normal continuous run", ""):
            rem += f" - {dt_reason}"

        # Match existing row by (date, shift, normalized POF)
        existing_match = existing_by_pof.get((s_date, s_st, s_norm_pof))
        if not existing_match and (s_date, s_st) in existing_by_date_shift:
            cand_list = existing_by_date_shift[(s_date, s_st)]
            if len(cand_list) == 1:
                existing_match = cand_list[0]

        if existing_match:
            # Row already exists in Excel! Check if update is needed
            ex_row = existing_match["row"]
            ex_good = existing_match["good"]
            ex_scrap = existing_match["scrap"]
            ex_dt = existing_match["dt"]

            # If already identical, skip
            if ex_good == s_good and ex_scrap == s_scrap and abs(ex_dt - s_dt) < 0.01:
                continue

            # Update row cells in Excel
            ws.cell(row=ex_row, column=col_date, value=s_date)
            if col_shift:
                ws.cell(row=ex_row, column=col_shift, value=s_st)
            ws.cell(row=ex_row, column=col_machine, value="Continuous Line 1")
            ws.cell(row=ex_row, column=col_pof, value=s_pof)
            ws.cell(row=ex_row, column=col_product, value=display_prod)
            ws.cell(row=ex_row, column=col_pid, value=pid)
            ws.cell(row=ex_row, column=col_customer, value=s["customer_name"] or "Alpha Standard")

            c_tot = ws.cell(row=ex_row, column=col_total, value=s_total)
            c_tot.number_format = '#,##0'

            c_good = ws.cell(row=ex_row, column=col_good, value=s_good)
            c_good.number_format = '#,##0'

            c_scrap = ws.cell(row=ex_row, column=col_scrap, value=s_scrap)
            c_scrap.number_format = '#,##0'

            c_rate = ws.cell(row=ex_row, column=col_scrap_pct, value=scrap_rate)
            c_rate.number_format = '0.00%'

            c_dt = ws.cell(row=ex_row, column=col_downtime, value=s_dt)
            c_dt.number_format = '0.00'

            ws.cell(row=ex_row, column=col_remarks, value=rem)

            existing_match["good"] = s_good
            existing_match["scrap"] = s_scrap
            existing_match["dt"] = s_dt
            rows_updated += 1
        else:
            # Append new row at the next available line
            target_row = max(last_data_row + 1, header_row_idx + 1)
            ws.cell(row=target_row, column=col_date, value=s_date)
            if col_shift:
                ws.cell(row=target_row, column=col_shift, value=s_st)
            ws.cell(row=target_row, column=col_machine, value="Continuous Line 1")
            ws.cell(row=target_row, column=col_pof, value=s_pof)
            ws.cell(row=target_row, column=col_product, value=display_prod)
            ws.cell(row=target_row, column=col_pid, value=pid)
            ws.cell(row=target_row, column=col_customer, value=s["customer_name"] or "Alpha Standard")

            c_tot = ws.cell(row=target_row, column=col_total, value=s_total)
            c_tot.number_format = '#,##0'

            c_good = ws.cell(row=target_row, column=col_good, value=s_good)
            c_good.number_format = '#,##0'

            c_scrap = ws.cell(row=target_row, column=col_scrap, value=s_scrap)
            c_scrap.number_format = '#,##0'

            c_rate = ws.cell(row=target_row, column=col_scrap_pct, value=scrap_rate)
            c_rate.number_format = '0.00%'

            c_dt = ws.cell(row=target_row, column=col_downtime, value=s_dt)
            c_dt.number_format = '0.00'

            ws.cell(row=target_row, column=col_remarks, value=rem)

            new_entry = {"row": target_row, "good": s_good, "scrap": s_scrap, "dt": s_dt, "norm_pof": s_norm_pof}
            if s_norm_pof:
                existing_by_pof[(s_date, s_st, s_norm_pof)] = new_entry
            if (s_date, s_st) not in existing_by_date_shift:
                existing_by_date_shift[(s_date, s_st)] = []
            existing_by_date_shift[(s_date, s_st)].append(new_entry)

            last_data_row = target_row
            rows_written += 1

    wb.save(excel_path)
    wb.close()

    return {
        "status": "success",
        "file": excel_path,
        "filename": os.path.basename(excel_path),
        "rows_written": rows_written,
        "rows_updated": rows_updated,
        "total_shifts_in_db": len(shifts),
        "message": f"Successfully updated {os.path.basename(excel_path)} ({rows_written} added, {rows_updated} updated)."
    }

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Alpha Aerosols Production Engine & Data Sync")
    parser.add_argument("--sync-excel", nargs="?", const="ACTIVE", default=None,
                        help="Sync production data from Excel workbook into SQLite database and update production.json")
    parser.add_argument("--sync-to-excel", nargs="?", const="ACTIVE", default=None,
                        help="Sync production shifts from SQLite database into Excel monthly workbook")
    parser.add_argument("--export-json", action="store_true", help="Re-export production.json snapshot from SQLite")
    parser.add_argument("--reseed", action="store_true", help="Reseed database with clean master baseline")
    parser.add_argument("--demo", action="store_true", help="Seed with demo historical shifts for testing")
    args = parser.parse_args()

    if args.sync_excel is not None:
        target_f = None if args.sync_excel == "ACTIVE" else args.sync_excel
        print(f"Syncing production records from Excel: {target_f or 'Auto-detected Monthly Workbook'}...")
        result = sync_excel_to_db(target_f)
        print(f"Excel -> DB Sync Complete: {result}")
    elif args.sync_to_excel is not None:
        target_f = None if args.sync_to_excel == "ACTIVE" else args.sync_to_excel
        print(f"Syncing shifts from DB to Excel: {target_f or 'Auto-detected Monthly Workbook'}...")
        result = sync_db_to_excel(target_f)
        print(f"DB -> Excel Sync Complete: {result}")
    elif args.reseed:
        print("Reseeding master database...")
        seed_master_data(force_reseed=True, demo_data=args.demo)
        export_production_json()
        print("Database reseeded successfully.")
    elif args.export_json:
        export_production_json()
        print(f"JSON snapshot updated at {JSON_PATH}")
    else:
        print("Initializing Alpha Aerosols Database Engine...")
        seed_master_data(demo_data=False)
        data = export_production_json()
        print(f"Data engine ready. Today output: {data['kpis']['today_output_cans']:,} cans.")
        print(f"JSON snapshot saved at: {JSON_PATH}")

